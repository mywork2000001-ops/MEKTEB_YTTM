"""Analitika: reytinq, səviyyə (avtomatik), risk, davamiyyət 25%+, şagird kartı, Excel, audit."""
import datetime as dt
import io

from app.models import PlanLesson, TeachingAssignment

SLOTS = {'0': [1], '1': [1], '2': [1], '3': [1], '4': [1]}        # hər gün 1 saat


def setup(world):
    as_, S = world
    c = as_('admin')
    cid = c.post('/api/classes', json={'name': 'X c'}).json()['id']
    c.post(f'/api/classes/{cid}/join', json={'subject': 'Riyaziyyat', 'weekly_hours': 5, 'slots': SLOTS})
    mk = lambda n, m: c.post('/api/students', json={'full_name': n, 'class_id': cid, 'score_math': m}).json()['id']
    good, weak, new = mk('Güclü Şagird oğlu', 35.0), mk('Zəif Şagird qızı', 80.0), mk('Yeni Şagird oğlu', 55.0)
    with S() as db:
        ta = db.query(TeachingAssignment).filter_by(class_id=cid).one()
        for i in range(40):
            db.add(PlanLesson(assignment_id=ta.id, seq=i + 1, semester=1, topic=f'M{i}', assessment_type='formativ',
                              date=dt.date(2026, 9, 15)))
        db.commit()
        ta_id = ta.id
    days = [dt.date(2026, 9, 15) + dt.timedelta(days=i) for i in range(14)]
    days = [d for d in days if d.weekday() < 5]
    for i, d in enumerate(days):
        c.put(f'/api/journal/{ta_id}/entry', json={
            'date': str(d), 'period': 1,
            'attendance': {good: 'var', weak: 'yox' if i % 2 == 0 else 'var', new: 'var'},
            'marks': [{'student_id': good, 'kind': 'şifahi', 'grade': 5}] +
                     ([{'student_id': weak, 'kind': 'şifahi', 'grade': 2}] if i % 2 else []),
            'homework_checks': {good: 'etdi', weak: 'etmədi'}})
    k = c.post(f'/api/exams/{ta_id}', json={'kind': 'KSQ', 'no': 1, 'semester': 1, 'date': '2026-09-25',
                                            'max_points': 20}).json()
    c.put(f'/api/exams/{ta_id}/{k["id"]}/scores', json=[{'student_id': good, 'points': 19},
                                                        {'student_id': weak, 'points': 4}])
    return c, ta_id, good, weak, new


def test_rating_levels_risk(world):
    c, ta, good, weak, new = setup(world)
    a = c.get(f'/api/analytics/{ta}', params={'semester': 1}).json()
    rows = {r['student_id']: r for r in a['students']}
    assert rows[good]['place'] == 1 and rows[weak]['place'] == 2 and rows[new]['place'] is None
    # səviyyə nəticələrə görə dəyişib: IX balı zəif olan güclü, güclü olan zəif
    assert rows[good]['baseline_level'] == 'Zəif' and rows[good]['level'] == 'Güclü'
    assert rows[weak]['baseline_level'] == 'Güclü' and rows[weak]['level'] == 'Zəif'
    assert rows[new]['level'] == 'Orta' and rows[new]['level_source'] == 'IX sinif balı'
    assert rows[weak]['risk']['status'] == 'Qırmızı'
    assert any('Davamiyyət' in f['izah'] and f['bal'] for f in rows[weak]['risk']['factors'])
    assert rows[weak]['absence_warning'] and not rows[good]['absence_warning']
    assert a['overview']['levels'] == {'Güclü': 1, 'Orta': 1, 'Zəif': 1}
    lv = c.get(f'/api/analytics/{ta}/levels').json()
    assert [x['full_name'] for x in lv['Zəif']] == ['Zəif Şagird qızı'] and lv['Məlum deyil'] == []


def test_attendance_heatmap(world):
    c, ta, good, weak, new = setup(world)
    m = c.get(f'/api/analytics/{ta}/attendance').json()
    assert len(m['columns']) == 10
    w = next(r for r in m['rows'] if r['student_id'] == weak)
    assert w['missed_pct'] == 50.0 and w['warning'] and w['cells'][0] == 'yox'


def test_student_card_private_to_teacher(world):
    as_, _ = world
    c, ta, good, weak, new = setup(world)
    ilqar = as_('ilqar')
    r = c.post(f'/api/students/{weak}/contacts', json={'date': '2026-09-26', 'method': 'zəng',
                                                        'topic': 'Davamiyyət', 'outcome': 'Valideyn məlumatlandırıldı'})
    assert r.status_code == 200
    p = c.post(f'/api/students/{weak}/plans', json={'goal': 'KSQ-2-də 50%+', 'start': '2026-09-26',
                                                     'steps': [{'text': 'Həftədə 2 əlavə dərs'}, {'text': 'Ev tapşırığı', 'done': True}]})
    assert p.json()['progress'] == 50
    assert ilqar.get(f'/api/students/{weak}/contacts').status_code == 403          # başqa müəllim görmür
    cid = c.get('/api/classes').json()[0]['id']
    ilqar.post(f'/api/classes/{cid}/join', json={'subject': 'Fizika', 'weekly_hours': 1, 'slots': {}})
    assert ilqar.get(f'/api/students/{weak}/contacts').json() == []               # qoşulsa da – yalnız öz qeydləri


def test_excel_export_and_audit(world):
    as_, _ = world
    c, ta, good, weak, new = setup(world)
    r = c.get(f'/api/reports/{ta}/xlsx', params={'semester': 1})
    assert r.status_code == 200 and r.headers['content-type'].startswith('application/vnd.openxml')
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws['B5'].value == 'Güclü Şagird oğlu' and ws.page_setup.orientation == 'portrait'
    log = c.get('/api/audit', params={'entity': 'exam_scores'}).json()
    assert log and log[0]['user'] == 'Həsənov Fərid'
    assert as_('ilqar').get('/api/audit').status_code == 403


def test_overview(world):
    c, ta, *_ = setup(world)
    ov = c.get('/api/overview').json()
    assert ov[0]['class_name'] == 'X c' and ov[0]['lessons_written'] == 10 and ov[0]['absence_warnings'] == 1


def test_manual_level_override(world):
    as_, _ = world
    c, ta, good, weak, new = setup(world)
    assert c.put(f'/api/analytics/{ta}/levels/{weak}', json={'level': 'Orta', 'note': 'dərsdə fəaldır'}).status_code == 200
    rows = {r['student_id']: r for r in c.get(f'/api/analytics/{ta}').json()['students']}
    assert rows[weak]['level'] == 'Orta' and rows[weak]['auto_level'] == 'Zəif' and rows[weak]['level_source'] == 'müəllim'
    lv = c.get(f'/api/analytics/{ta}/levels').json()
    assert weak in [x['student_id'] for x in lv['Orta']]
    c.put(f'/api/analytics/{ta}/levels/{weak}', json={'level': None})
    rows = {r['student_id']: r for r in c.get(f'/api/analytics/{ta}').json()['students']}
    assert rows[weak]['level'] == 'Zəif' and rows[weak]['manual_level'] is None
    assert as_('ilqar').put(f'/api/analytics/{ta}/levels/{weak}', json={'level': 'Güclü'}).status_code == 404


def test_journal_methodology_rules(world, monkeypatch):
    """Metodist: üzrlüyə qiymət yox; gələcəyə yalnız mövzu/ev tapşırığı; qayıbın ev tapşırığı yoxlanmır;
    eyni gün verilən ev tapşırığı həmin gün yoxlanmır; KSQ tarixi yarımil içində; yarımildə bir BSQ."""
    monkeypatch.setattr('app.api.journal.today', lambda: dt.date(2026, 9, 29))
    c, ta, good, weak, new = setup(world)
    put = lambda **b: c.put(f'/api/journal/{ta}/entry', json=b)
    assert put(date='2026-09-29', period=1, attendance={weak: 'üzrlü'},
               marks=[{'student_id': weak, 'kind': 'şifahi', 'grade': 4}]).status_code == 400
    assert put(date='2026-10-01', period=1, attendance={good: 'var'}).status_code == 400
    assert put(date='2026-10-01', period=1, homework='S 1–5').status_code == 200
    r = put(date='2026-09-29', period=1, homework='E 3', attendance={weak: 'yox', good: 'var'},
            homework_checks={weak: 'etmədi', good: 'etdi'}).json()
    assert r['homework_checks'] == {str(good): 'etdi'}
    with world[1]() as db:                                   # ikinci saat həmin gün – ev tapşırığı hələ yoxlanmır
        tao = db.get(TeachingAssignment, ta)
        tao.slots = {**tao.slots, '1': [1, 5]}
        db.commit()
    day = c.get(f'/api/journal/{ta}/day', params={'date': '2026-09-29'}).json()['lessons']
    assert day[1]['homework_to_check'] != 'E 3'
    nxt = c.get(f'/api/journal/{ta}/day', params={'date': '2026-09-30'}).json()['lessons']
    assert nxt[0]['homework_to_check'] == 'E 3'
    ex = lambda **b: c.post(f'/api/exams/{ta}', json={'max_points': 20, **b}).status_code
    assert ex(kind='KSQ', no=5, semester=1, date='2027-03-01') == 400
    assert ex(kind='BSQ', no=2, semester=1, date='2027-01-20') == 400
    assert ex(kind='BSQ', no=1, semester=1, date='2027-01-20') == 200
    assert ex(kind='BSQ', no=1, semester=1, date='2027-01-21') == 409


def test_journal_grid_plan_execution_and_feedback(world, monkeypatch):
    """Jurnal səhifəsi (ay), planın icrası, formativ rəy və dərs qeydi, yazmayan şagird KSQ-də «2» almır."""
    monkeypatch.setattr('app.api.journal.today', lambda: dt.date(2026, 9, 30))
    monkeypatch.setattr('app.api.plan.today', lambda: dt.date(2026, 10, 1))      # 30.09 yazılmayıb – gecikir
    c, ta, good, weak, new = setup(world)
    r = c.put(f'/api/journal/{ta}/entry', json={'date': '2026-09-29', 'period': 1, 'note': 'Qrup işi',
                                                'attendance': {good: 'var', weak: 'gecikdi', new: 'var'},
                                                'marks': [{'student_id': good, 'kind': 'şifahi', 'grade': 5,
                                                           'comment': 'Həlli əsaslandırdı'}]}).json()
    assert r['note'] == 'Qrup işi' and r['marks'][0]['comment'] == 'Həlli əsaslandırdı'
    g = c.get(f'/api/journal/{ta}/grid', params={'month': '2026-09'}).json()
    assert len(g['columns']) == 12 and g['columns'][-1]['written'] is False      # 30.09 yazılmayıb
    rows = {x['student_id']: x for x in g['rows']}
    assert rows[weak]['cells'][0]['att'] == 'q' and rows[weak]['cells'][-2]['att'] == 'g'
    assert rows[weak]['missed'] == 5 and rows[good]['cells'][-2]['marks'] == [5]
    assert c.get(f'/api/journal/{ta}/grid', params={'month': 'x'}).status_code == 400
    off = c.get(f'/api/plan/{ta}/official').json()
    st = [x['status'] for x in off]
    assert st.count('keçilib') == 11 and st.count('gecikir') == 1 and off[0]['taught_dates'] == ['2026-09-15']
    # KSQ: tapşırıq qeydi olmayan şagird – «daxil edilməyib», qiymət yoxdur (0 bal/«2» deyil)
    k = c.post(f'/api/exams/{ta}', json={'kind': 'KSQ', 'no': 2, 'semester': 1, 'date': '2026-10-01',
                                         'items': [{'n': 1, 'points': 2, 'standard': '1.2.5'}, {'n': 2, 'points': 2}]}).json()
    res = c.put(f'/api/exams/{ta}/{k["id"]}/scores', json=[{'student_id': good, 'item_marks': [1, 1]},
                                                          {'student_id': new, 'item_marks': None}]).json()
    rr = {x['student_id']: x for x in res['rows']}
    assert rr[good]['grade'] == 5 and rr[new]['grade'] is None and rr[new]['points'] is None
    assert res['standards'] == [{'standard': '1.2.5', 'pct': 100.0}]
