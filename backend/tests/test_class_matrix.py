"""Sinif matrisi «kim nəyi işləyib» (docs/tasks/sinif-test-matrisi.md): statuslar, təyin olunmayan, açıq test, kağız KSQ,
Excel/PDF, başqa müəllim – 404."""
import datetime as dt
import io

import pytest
from fastapi import HTTPException

from app.models import OnlineTask, PlanLesson
from app.student_report import class_matrix

from .test_api_portal import UTC, clock, setup  # noqa: F401
from .test_student_report import WIN, _mk, _paper, _staff, patched  # noqa: F401


@pytest.fixture
def mat(world, patched, monkeypatch):
    """3 şagird: mövzu testi A (2.3; yalnız 0 və 1-ci şagirdə, 0 yazıb), mövzu testi B (hələ açıq), sınaq ÜSİ-1 (0 yazıb),
    kağız KSQ-1 (0-ın balı nəşr olunmayıb)."""
    as_, S = world
    admin, ta, st, bid = _mk(world, patched)
    monkeypatch.setattr('app.services.today', lambda: dt.date(2026, 10, 10))
    with S() as db:
        a = db.query(OnlineTask).filter_by(kind='movzu').one()
        a.student_ids = [st[0]['id'], st[1]['id']]
        pl3 = db.query(PlanLesson).filter_by(assignment_id=ta, seq=3).one().id
        db.commit()
    win = {'opens_at': '2026-09-29T20:00:00Z', 'closes_at': '2026-10-02T14:00:00Z'}
    r = admin.post(f'/api/plan/{ta}/topics/{pl3}/test', json={'title': 'Mövzu B – test', 'duration_min': 20, 'bank_ids': _ids(S, ta),
                                                              'targets': [{'ta_id': ta, 'plan_lesson_id': pl3, **win}]})
    assert r.status_code == 200, r.text
    _paper(S, ta, st[0]['id'], 'KSQ', 1, 24, 30, published=False, scan=True)
    return admin, S, ta, st


def _ids(S, ta):
    from app.models import BankQuestion
    with S() as db:
        return [q.id for q in db.query(BankQuestion).order_by(BankQuestion.id)]


def test_matrix_statuses(mat):
    admin, S, ta, st = mat
    with S() as db:
        m = class_matrix(db, _staff(S), ta, 'all')
    cols = {c['key']: c for c in m['columns']}
    assert [c['kind'] for c in m['columns']] == ['movzu', 'movzu', 'sinaq', 'ksq']
    b, a, sq, ksq = m['columns']                        # mövzu testləri № sırası ilə: B (2.1), A (2.3)
    assert a['no'] == '2.3' and b['no'] == '2.1' and sq['title'] == 'ÜSİ-1' and ksq['title'] == 'KSQ-1'
    rows = {r['student_id']: r for r in m['rows']}
    s0, s1, s2 = (rows[x['id']] for x in st)
    # mövzu testi A: 0 yazıb (100%), 1 yazmayıb, 2 təyin olunmayıb
    assert (s0['cells'][a['key']]['status'], s0['cells'][a['key']]['pct'], s0['cells'][a['key']]['grade']) == ('yazıb', 100.0, 5)
    assert s1['cells'][a['key']]['status'] == 'yazmayıb' and s2['cells'][a['key']]['status'] == 'təyin olunmayıb'
    assert a['assigned_count'] == 2 and a['wrote_count'] == 1 and a['avg_pct'] == 100.0
    # açıq test B – hamı ⏳
    assert {r['cells'][b['key']]['status'] for r in rows.values()} == {'gözlənilir'}
    # sınaq: 0 yazıb, qalanı yazmayıb
    assert s0['cells'][sq['key']]['status'] == 'yazıb' and s2['cells'][sq['key']]['status'] == 'yazmayıb'
    # kağız KSQ: nəşr olunmayıb (bal var, işarəli), qalanı yazmayıb
    k0 = s0['cells'][ksq['key']]
    assert (k0['status'], k0['pct'], k0['published']) == ('nəşr olunmayıb', 80.0, False)
    assert s1['cells'][ksq['key']]['status'] == 'yazmayıb'
    # sətir yekunu: 0 – 3 yazıb (A, ÜSİ, KSQ-nəşrsiz), verilən 3; orta yalnız nəşr olunanlardan
    assert (s0['wrote'], s0['given'], s0['avg_pct']) == (3, 3, 100.0)
    assert (s2['wrote'], s2['given']) == (0, 2)
    assert cols[ksq['key']]['assigned_count'] == 3


def test_matrix_kind_filter_and_period(mat):
    admin, S, ta, st = mat
    with S() as db:
        assert [c['kind'] for c in class_matrix(db, _staff(S), ta, 'movzu')['columns']] == ['movzu', 'movzu']
        assert [c['title'] for c in class_matrix(db, _staff(S), ta, 'sinaq')['columns']] == ['ÜSİ-1']
        assert class_matrix(db, _staff(S), ta, 'all', dt.date(2026, 12, 1))['columns'] == []
        with pytest.raises(HTTPException) as e:
            class_matrix(db, _staff(S), ta, 'x')
        assert e.value.status_code == 400


def test_matrix_other_teacher_404(mat):
    admin, S, ta, st = mat
    with S() as db:
        with pytest.raises(HTTPException) as e:
            class_matrix(db, _staff(S, 'yad'), ta, 'all')
        assert e.value.status_code == 404


def test_matrix_api_and_exports(mat, world):
    admin, S, ta, st = mat
    as_, _ = world
    j = admin.get(f'/api/results-center/class/{ta}/matrix', params={'kind': 'all'})
    assert j.status_code == 200, j.text
    assert len(j.json()['columns']) == 4 and len(j.json()['rows']) == 3
    x = admin.get(f'/api/results-center/class/{ta}/matrix', params={'format': 'xlsx'})
    assert x.status_code == 200 and x.content[:2] == b'PK'
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.cell(5, 2).value == 'Şagird' and any('(80)' == str(c.value) for row in ws.iter_rows() for c in row)
    assert as_('yad').get(f'/api/results-center/class/{ta}/matrix').status_code == 404
    assert as_('yad').get(f'/api/results-center/class/{ta}/matrix', params={'format': 'xlsx'}).status_code == 404


def test_matrix_html_print(mat):
    from app.class_matrix_print import matrix_html
    admin, S, ta, st = mat
    with S() as db:
        m = class_matrix(db, _staff(S), ta, 'all')
    h = matrix_html(m, 'Məktəb', 'Müəllim')
    assert 'A4 landscape' in h and '(80)' in h and 'ÜSİ-1' in h and 'gözlənilir' in h
