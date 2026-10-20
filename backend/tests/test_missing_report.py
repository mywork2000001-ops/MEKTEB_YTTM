"""İşlənməyənlər hesabatı (docs/tasks/islenmeyenler-hesabati.md): göstəricilər, risk, plan günü, təkrarsızlıq, H1–H4,
«Tələb et», «Araşdır», səbəb qeydləri; köhnə /missing eyni saylar."""
import datetime as dt

import pytest

from app.missing_list_print import all_classes_html, class_html, request_text, share_text, short_name, students_html
from app.models import MissingNote, MissingRequest, OnlineTask
from app.student_report import MISSING_RISK, class_matrix, class_missing, plan_state, risk_of
from app.timing import NO_DEADLINE

from .test_api_portal import UTC, clock, setup  # noqa: F401
from .test_class_matrix import mat  # noqa: F401
from .test_student_report import WIN, _mk, _paper, _staff, patched  # noqa: F401

U = '/api/results-center'


def _rows(S, ta, **kw):
    with S() as db:
        m = class_missing(db, _staff(S), ta, 'all', all_students=True, **kw)
    return m, {r['student_id']: r for r in m['rows']}


def test_risk_thresholds_and_streak():
    assert MISSING_RISK == {'critical': 50, 'attention': 80, 'streak': 3}
    assert risk_of(49.9, 0) == 'critical' and risk_of(50.0, 0) == 'attention' and risk_of(79.9, 0) == 'attention'
    assert risk_of(80.0, 0) == 'ok' and risk_of(None, 0) == 'ok'
    assert risk_of(100.0, 2) == 'ok' and risk_of(100.0, 3) == 'critical'          # ardıcıl ≥ 3 buraxma


def test_plan_state_rules():
    col = {'plan_date': '2026-10-05'}
    td = dt.date(2026, 10, 10)
    assert plan_state('təyin olunmayıb', col, {}, td) == (None, None)
    assert plan_state('yazmayıb', col, {}, td) == ('late_open', 5)
    assert plan_state('gözlənilir', col, {}, td) == ('late_open', 5)
    assert plan_state('gözlənilir', {'plan_date': '2026-10-12'}, {}, td) == ('not_due', None)
    assert plan_state('gözlənilir', {'plan_date': '2026-10-10'}, {}, td) == ('not_due', None)      # plan günü bu gün – hələ borc deyil
    assert plan_state('yazıb', col, {'submitted_at': '2026-10-08T10:00:00+00:00'}, td) == ('late_done', 3)
    assert plan_state('yazıb', col, {'submitted_at': '2026-10-05T10:00:00+00:00'}, td) == ('on_time', None)
    assert plan_state('nəşr olunmayıb', col, {}, td) == ('on_time', None)


def test_metrics_and_summary(mat):
    admin, S, ta, st = mat
    m, rows = _rows(S, ta)
    s0, s1, s2 = (rows[x['id']] for x in st)
    assert (s0['assigned'], s0['done'], s0['missed'], s0['pending'], s0['rate']) == (4, 3, 0, 1, 75.0)
    assert (s1['assigned'], s1['done'], s1['missed'], s1['pending'], s1['rate']) == (4, 0, 3, 1, 0.0)
    assert (s2['assigned'], s2['done'], s2['missed'], s2['pending']) == (3, 0, 2, 1)
    assert s1['risk'] == 'critical' and s1['advice_key'] == 'contact' and s1['streak'] >= 1
    assert {r['risk'] for r in m['rows']} <= {'critical', 'attention', 'ok'}
    sm = m['summary']
    assert sm['students_total'] == 3 and sm['tests'] == 4 and sm['full'] == 0 and sm['critical'] >= 2
    assert sm['avg_rate'] == round((75.0 + 0.0 + 0.0) / 3, 1)
    assert len(m['tests']) == 4 and {'key', 'date', 'title', 'no', 'assigned', 'done', 'rate', 'missed'} <= set(m['tests'][0])
    t = {x['title']: x for x in m['tests']}['ÜSİ-1']
    assert (t['assigned'], t['done'], t['missed'], t['rate']) == (3, 1, 2, 33.3)
    # sıralama: risk → icra % → ad
    with S() as db:
        rk = class_missing(db, _staff(S), ta, 'all', all_students=True, sort='risk')['rows']
    order = {'critical': 0, 'attention': 1, 'ok': 2}
    assert [order[r['risk']] for r in rk] == sorted(order[r['risk']] for r in rk)
    assert rk[-1]['student_id'] == st[0]['id']


def test_default_rows_all_and_only_missed(mat):
    admin, S, ta, st = mat
    with S() as db:
        dflt = class_missing(db, _staff(S), ta, 'all')
        only = class_missing(db, _staff(S), ta, 'all', only_missed=True)
        every = class_missing(db, _staff(S), ta, 'sinaq', all_students=True)
    assert dflt['students'] == 3 and dflt['total'] == 8
    assert {r['student_id'] for r in only['rows']} == {st[1]['id'], st[2]['id']}       # s0-ın yalnız «gözlənilir»i var
    assert all(i['status'] == 'yazmayıb' for r in only['rows'] for i in r['items']) and only['total'] == 5
    assert every['students'] == 3                                                         # hamı (s0 sınaqda borcsuzdur)


def test_plan_dates_and_late_filter(mat):
    admin, S, ta, st = mat
    with S() as db:
        mx = class_matrix(db, _staff(S), ta, 'all')
        sq = next(c for c in mx['columns'] if c['kind'] == 'sinaq')
        task = db.query(OnlineTask).filter_by(kind='sinaq').first()
        from app.test_stats import local_date
        assert sq['plan_date'] == local_date(task.closes_at).isoformat()                 # plan dərsi yoxdur → son tarix
        a = next(c for c in mx['columns'] if c['no'] == '2.3')
        assert a['plan_date'] is not None                                                 # plan dərsinin tarixi
        task.closes_at, task.duration_min = NO_DEADLINE, 0                                # müddətsiz → açılış + 7 gün
        db.commit()
        mx2 = class_matrix(db, _staff(S), ta, 'all')
        sq2 = next(c for c in mx2['columns'] if c['kind'] == 'sinaq')
        assert sq2['plan_date'] == (local_date(task.opens_at) + dt.timedelta(days=7)).isoformat()
    m, rows = _rows(S, ta, plan='late_open')
    for r in m['rows']:
        assert all(i['state'] == 'late_open' for i in r['items'])
    full, _ = _rows(S, ta)
    assert m['total'] <= full['total']
    s1 = rows[st[1]['id']]
    assert s1['late_open'] >= 1 and s1['late_days'] >= 0 and 'rate_on_time' in s1 and s1['rate_total'] == s1['rate']
    assert all('plan_date' in i and 'late_days' in i for i in s1['items'])


def test_late_done_and_on_time_rate(mat):
    admin, S, ta, st = mat
    m, rows = _rows(S, ta)
    s0 = rows[st[0]['id']]
    assert isinstance(s0['late_done'], list) and s0['rate_on_time'] is not None


def test_unique_rows_and_columns(mat, monkeypatch):
    admin, S, ta, st = mat
    from app import student_report
    orig = student_report.roster
    monkeypatch.setattr('app.student_report.roster', lambda db, t, since=None: orig(db, t, since) * 2)
    m, rows = _rows(S, ta)
    assert len(m['rows']) == 3 and len(rows) == 3 and m['summary']['students_total'] == 3
    with S() as db:
        assert len(class_matrix(db, _staff(S), ta, 'all')['rows']) == 3


def test_resend_copy_is_one_column(mat):
    admin, S, ta, st = mat
    before, _ = _rows(S, ta)
    sq = next(t for t in before['tests'] if t['kind'] == 'sinaq')
    r = admin.post(f'/api/tasks/{ta}/{sq["task_id"]}/copy', json={'target_ta_id': ta, 'student_ids': [st[2]['id']],
                                                                   'opens_at': '2026-09-29T20:00:00Z', 'closes_at': '2026-10-30T12:00:00Z'})
    assert r.status_code == 200, r.text
    after, rows = _rows(S, ta)
    assert len(after['tests']) == len(before['tests']) == 4 and len(after['rows']) == 3
    assert [t['key'] for t in after['tests']] == [t['key'] for t in before['tests']]
    assert rows[st[2]['id']]['assigned'] == 3


def test_notes_excused_leave_denominator(mat):
    admin, S, ta, st = mat
    _, rows = _rows(S, ta)
    sq = next(t for t in _rows(S, ta)[0]['tests'] if t['kind'] == 'sinaq')['key']
    base = rows[st[2]['id']]
    body = {'task_key': sq, 'student_id': st[2]['id'], 'reason': 'excused', 'note': 'xəstə'}
    assert admin.put(f'{U}/class/{ta}/missing-note', json=body).status_code == 200
    m, rows = _rows(S, ta)
    x = rows[st[2]['id']]
    assert x['assigned'] == base['assigned'] - 1 and len(x['excused']) == 1 and x['excused'][0]['note'] == 'xəstə'
    assert all(i['key'] != sq for i in x['items'])
    # səbəb (üzrsüz) – qalır, amma səbəb göstərilir
    assert admin.put(f'{U}/class/{ta}/missing-note', json={**body, 'reason': 'forgot'}).status_code == 200
    m, rows = _rows(S, ta)
    it = next(i for i in rows[st[2]['id']]['items'] if i['key'] == sq)
    assert it['reason'] == 'forgot' and rows[st[2]['id']]['reasons'] == {'forgot': 1} and m['summary']['reasons'] == {'forgot': 1}
    with S() as db:
        assert db.query(MissingNote).count() == 1                                         # bir test + bir şagird = bir qeyd
    assert admin.put(f'{U}/class/{ta}/missing-note', json={'task_key': sq, 'student_id': st[2]['id']}).status_code == 200
    with S() as db:
        assert db.query(MissingNote).count() == 0
    assert admin.put(f'{U}/class/{ta}/missing-note', json={**body, 'task_key': 'zzz'}).status_code == 404
    assert admin.put(f'{U}/class/{ta}/missing-note', json={**body, 'student_id': 999999}).status_code == 404
    assert admin.put(f'{U}/class/{ta}/missing-note', json={**body, 'reason': 'x'}).status_code == 422


def test_share_text_name_modes(mat):
    admin, S, ta, st = mat
    with S() as db:
        m = class_missing(db, _staff(S), ta, 'all', sort='risk')
    full = share_text(m, 'az', 'full')
    short = share_text(m, 'az', 'short')
    none = share_text(m, 'az', 'none')
    names = [r['full_name'] for r in m['rows']]
    assert all(n in full for n in names)
    assert not any(n in short for n in names) and all(short_name(n) in short for n in names)
    assert not any(n in none for n in names) and not any(short_name(n) in none for n in names) and 'ÜSİ-1' in none
    assert 'Среднее выполнение' in share_text(m, 'ru', 'none') and 'Average completion' in share_text(m, 'en', 'none')
    assert short_name('Əliyev Elvin Rəşad') == 'Elvin Ə.'
    api = admin.get(f'{U}/class/{ta}/missing-share', params={'lang': 'en', 'names': 'short'})
    assert api.status_code == 200 and 'Average completion' in api.json()['text']
    assert admin.get(f'{U}/class/{ta}/missing-share', params={'names': 'x'}).status_code == 422
    empty = share_text({**m, 'rows': []}, 'az', 'full')
    assert 'bütün şagirdlər' in empty


def test_html_forms(mat):
    admin, S, ta, st = mat
    m, _ = _rows(S, ta, sort='risk')
    h = class_html(m, 'Məktəb', 'Müəllim')
    for s in ('Cədvəl 1', 'Cədvəl 2', 'Şərti işarələr', 'Kritik', 'Gecikmə (gün)', 'imza', 'ÜSİ-1'):
        assert s in h
    s = students_html(m, 'Məktəb', 'Müəllim')
    assert 'Valideyn' in s and 'Müəllim qeydi' in s
    all_h = all_classes_html([m, m], 'Məktəb', 'Müəllim', '01.10.2026 – 31.10.2026')
    assert all_h.count('Cədvəl 1') == 2 and 'Bütün siniflərim' in all_h


def test_api_pdf_xlsx_all_classes(mat, world):
    admin, S, ta, st = mat
    as_, _ = world
    u = f'{U}/class/{ta}/missing-list'
    j = admin.get(u, params={'all': 1, 'sort': 'risk'}).json()
    assert j['students'] == 3 and j['rows'][0]['risk'] == 'critical' and j['summary']['students_total'] == 3
    assert admin.get(u, params={'only_missed': 1}).json()['students'] == 2
    assert admin.get(u, params={'plan': 'bad'}).status_code == 422
    p = admin.get(u, params={'format': 'pdf'})
    assert p.status_code == 200 and p.content[:4] == b'%PDF'
    import io
    from openpyxl import load_workbook
    x = admin.get(u, params={'format': 'xlsx'})
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ['Şagirdlər', 'Testlər']
    assert wb['Şagirdlər'].cell(5, 2).value == 'Şagird' and wb['Testlər'].cell(1, 3).value == 'Test'
    assert wb['Testlər'].max_row == 1 + 4
    a = admin.get(f'{U}/missing-all', params={'format': 'pdf'})
    assert a.status_code == 200 and a.content[:4] == b'%PDF'
    assert admin.get(f'{U}/missing-all', params={'format': 'json'}).json()['classes'][0]['summary']['students_total'] == 3
    assert as_('yad').get(f'{U}/missing-all').status_code == 200                        # öz sinifləri yoxdur – boş xülasə


def test_old_missing_endpoint_matches_report(mat):
    """İkinci hesablama yoxdur: /missing və class_missing eyni şagird və say."""
    admin, S, ta, st = mat
    old = admin.get(f'{U}/missing', params={'ta_id': ta}).json()
    new = admin.get(f'{U}/class/{ta}/missing-list', params={'all': 1}).json()
    exp = {r['student_id']: sum(1 for i in r['items'] if i['status'] == 'yazmayıb') for r in new['rows'] if r['missed']}
    assert {x['student_id']: x['missed'] for x in old['students']} == exp
    assert sum(t['missed'] for t in old['tests']) == sum(exp.values())
    again = admin.get(f'{U}/missing').json()                                              # ta_id verilmədən (hamısı)
    assert {x['student_id']: x['missed'] for x in again['students']} == exp


def test_request_notifies_and_history(mat):
    admin, S, ta, st = mat
    r = admin.post(f'{U}/class/{ta}/missing-request', json={'student_ids': [st[1]['id'], st[2]['id']]})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out['sent'] == 2 and out['delivered'] + out['no_account'] == 2 and out['reopened'] == 0
    with S() as db:
        rq = db.query(MissingRequest).all()
        assert len(rq) == 2 and all(q.task_keys for q in rq)
        if out['delivered']:
            from app.models import ChatMessage
            assert db.query(ChatMessage).count() >= out['delivered']
    admin.post(f'{U}/class/{ta}/missing-request', json={'student_ids': [st[1]['id']], 'message': 'Zəhmət olmasa tamamla'})
    _, rows = _rows(S, ta)
    assert rows[st[1]['id']]['requests']['count'] == 2 and rows[st[1]['id']]['requests']['last']
    assert rows[st[0]['id']]['requests']['count'] == 0
    # səhv: yad şagird, boş siyahı, tarixsiz reopen
    bad = admin.post(f'{U}/class/{ta}/missing-request', json={'student_ids': [999999]}).json()
    assert bad['sent'] == 0 and bad['errors']
    assert admin.post(f'{U}/class/{ta}/missing-request', json={'student_ids': []}).status_code == 422
    assert admin.post(f'{U}/class/{ta}/missing-request', json={'student_ids': [st[1]['id']], 'reopen': True}).status_code == 400


def test_request_text_languages():
    items = [{'no': '2.3', 'title': 'Test'}]
    az, en, ru = (request_text(lg, 'Ad', items, '30.10.2026 10:00') for lg in ('az', 'en', 'ru'))
    assert 'Salam' in az and 'Hello' in en and 'Здравствуйте' in ru and '2.3 Test' in ru and '30.10.2026' in en
    assert request_text('az', 'Ad', items, None, 'Mənim mətnim').startswith('Mənim mətnim')


def test_request_reopen_creates_copy_for_selected_only(mat):
    admin, S, ta, st = mat
    before = admin.get(f'{U}/class/{ta}/missing-list', params={'all': 1}).json()
    sq = next(t for t in before['tests'] if t['kind'] == 'sinaq')
    with S() as db:
        n0 = db.query(OnlineTask).count()
    r = admin.post(f'{U}/class/{ta}/missing-request', json={'student_ids': [st[2]['id']], 'task_keys': [sq['key']],
                                                            'reopen': True, 'new_deadline': '2026-10-30T12:00:00Z'})
    assert r.status_code == 200, r.text
    assert r.json()['reopened'] == 1
    with S() as db:
        assert db.query(OnlineTask).count() == n0 + 1
        cp = db.query(OnlineTask).order_by(OnlineTask.id.desc()).first()
        assert cp.student_ids == [st[2]['id']]
        assert db.query(MissingRequest).one().reopened is True
    after = admin.get(f'{U}/class/{ta}/missing-list', params={'all': 1}).json()
    assert len(after['tests']) == len(before['tests'])                                   # surət ayrı sütun yaratmır


def test_investigate_facts(mat):
    admin, S, ta, st = mat
    admin.post(f'{U}/class/{ta}/missing-request', json={'student_ids': [st[1]['id']]})
    r = admin.get(f"{U}/student/{st[1]['id']}/missing-investigate", params={'ta': ta})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d['student']['id'] == st[1]['id'] and d['risk'] == 'critical' and len(d['items']) == 4
    f = d['items'][0]['facts']
    assert {'visible', 'attempt_started', 'requests', 'attendance_on_plan_day', 'last_login', 'similar_done'} <= set(f)
    assert f['requests']['count'] == 1 and f['last_login'] is None
    assert len(d['history']) == 1 and d['history'][0]['reopened'] is False
    assert admin.get(f'{U}/student/999999/missing-investigate', params={'ta': ta}).status_code == 404
