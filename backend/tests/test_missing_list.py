"""İşlənməyənlər siyahısı (docs/tasks/islenmeyenler-siyahisi.md): matrisdən qurulur; sinif/şagird çapı, Excel, 404, Əsas səhifə sayı."""
import io

import pytest
from fastapi import HTTPException

from app.student_report import MISSING_STATUSES, class_matrix, class_missing

from .test_api_portal import UTC, clock, setup  # noqa: F401
from .test_class_matrix import mat  # noqa: F401
from .test_student_report import WIN, _mk, _paper, _staff, patched  # noqa: F401


def test_missing_statuses_constant():
    assert 'yazmayıb' in MISSING_STATUSES and 'gözlənilir' in MISSING_STATUSES and 'yazıb' not in MISSING_STATUSES


def test_missing_list_content(mat):
    admin, S, ta, st = mat
    with S() as db:
        m = class_missing(db, _staff(S), ta, 'all')
        rows = {r['student_id']: r for r in m['rows']}
        # s0: yalnız açıq mövzu testi B (qalanını yazıb / KSQ balı var); s1: A, B, sınaq, KSQ; s2: B, sınaq, KSQ (A ona təyin olunmayıb)
        s0, s1, s2 = (rows[x['id']] for x in st)
        assert [(i['kind'], i['status']) for i in s0['items']] == [('movzu', 'gözlənilir')]
        assert len(s1['items']) == 4 and len(s2['items']) == 3
        assert {i['status'] for i in s2['items']} == {'yazmayıb', 'gözlənilir'}
        assert m['students'] == 3 and m['total'] == 8
        # hamısını yazan şagird siyahıda yoxdur (yalnız sınaq: s0 yazıb)
        sq = {r['student_id'] for r in class_missing(db, _staff(S), ta, 'sinaq')['rows']}
        assert sq == {st[1]['id'], st[2]['id']}
        # hər element matrisdəki xana ilə üst-üstə düşür
        mx = class_matrix(db, _staff(S), ta, 'all')
        for r in m['rows']:
            mr = next(x for x in mx['rows'] if x['student_id'] == r['student_id'])
            assert {i['key'] for i in r['items']} == {k for k, c in mr['cells'].items() if c['status'] in MISSING_STATUSES}
        # sid: boş siyahılı şagird də qalır
        one = class_missing(db, _staff(S), ta, 'sinaq', sid=st[0]['id'])
        assert one['students'] == 1 and one['rows'][0]['items'] == []
        assert [r['full_name'] for r in m['rows']] == sorted(r['full_name'] for r in m['rows'])


def test_missing_other_teacher_404(mat):
    admin, S, ta, st = mat
    with S() as db:
        with pytest.raises(HTTPException) as e:
            class_missing(db, _staff(S, 'yad'), ta, 'all')
        assert e.value.status_code == 404


def test_missing_api_exports(mat, world):
    admin, S, ta, st = mat
    as_, _ = world
    u = f'/api/results-center/class/{ta}/missing-list'
    j = admin.get(u)
    assert j.status_code == 200 and j.json()['total'] == 8
    p = admin.get(u, params={'format': 'pdf'})
    assert p.status_code == 200 and p.content[:4] == b'%PDF'
    p1 = admin.get(u, params={'format': 'pdf', 'per_student': 1})
    assert p1.status_code == 200 and p1.content[:4] == b'%PDF'
    p2 = admin.get(u, params={'format': 'pdf', 'per_student': 1, 'sid': st[1]['id']})
    assert p2.status_code == 200 and p2.content[:4] == b'%PDF'
    x = admin.get(u, params={'format': 'xlsx'})
    assert x.status_code == 200 and x.content[:2] == b'PK'
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.cell(5, 2).value == 'Şagird' and sum(1 for r in ws.iter_rows(min_row=6) if r[1].value and str(r[1].value).startswith('Cəmi')) == 1
    assert admin.get(u, params={'format': 'xlsx', 'per_student': 1}).status_code == 400
    assert admin.get(u, params={'sid': 999999}).status_code == 404
    for params in ({}, {'format': 'pdf'}, {'format': 'xlsx'}, {'format': 'pdf', 'per_student': 1}):
        assert as_('yad').get(u, params=params).status_code == 404


def test_missing_student_sheet(mat, world):
    admin, S, ta, st = mat
    as_, _ = world
    p = admin.get(f"/api/results-center/student/{st[1]['id']}/missing-list")
    assert p.status_code == 200 and p.content[:4] == b'%PDF'
    assert as_('yad').get(f"/api/results-center/student/{st[1]['id']}/missing-list").status_code == 404


def test_missing_html_languages(mat):
    from app.missing_list_print import class_html, students_html
    admin, S, ta, st = mat
    with S() as db:
        m = class_missing(db, _staff(S), ta, 'all')
    h = class_html(m, 'Məktəb', 'Müəllim')
    assert 'Mövzu testləri' in h and 'Sınaqlar' in h and 'ÜSİ-1' in h
    m['rows'][1]['lang'] = 'ru'
    s = students_html(m, 'Məktəb', 'Müəllim')
    assert s.count('page-break-before') == len(m['rows']) - 1
    assert s.count('Тесты, которые нужно выполнить') == 1


def test_missing_summary_matches_list(mat):
    admin, S, ta, st = mat
    s = admin.get('/api/results-center/missing-summary')
    assert s.status_code == 200
    items = s.json()['items']
    lst = admin.get(f'/api/results-center/class/{ta}/missing-list').json()
    assert len(items) == 1 and items[0]['ta_id'] == ta
    assert (items[0]['students'], items[0]['total']) == (lst['students'], lst['total'])
