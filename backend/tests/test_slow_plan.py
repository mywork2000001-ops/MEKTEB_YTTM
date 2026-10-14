"""Zəif qrup tempi (docs/x-e-zeif-qrup-tempi-promtu.md)."""
import datetime as dt

from app.models import PlanLesson, PlanProgram, TeachingAssignment

from .test_api_portal import setup


def _plan(S, ta, n1=60, n2=50):
    with S() as db:
        db.query(PlanLesson).filter_by(assignment_id=ta).delete()
        seq = 0
        for sem, n in ((1, n1), (2, n2)):
            for i in range(n):
                seq += 1
                sec = f'{"I" if sem == 1 else "XIX"} BÖLMƏ – B{sem}{i // 10}'
                typ = 'KSQ' if i % 12 == 11 else 'formativ'
                db.add(PlanLesson(assignment_id=ta, seq=seq, semester=sem, section=sec, topic=f'Mövzu {sem}-{i}', standards=[],
                                  assessment_type=typ, exam_no=(i // 12 + 1) if typ == 'KSQ' else None, date=dt.date(2026, 9, 15), tasks=[]))
        db.commit()


def test_slow_plan_preview_and_apply(world):
    as_, S = world
    admin, ta, cid, st, _ = setup(world)
    _plan(S, ta, n1=24, n2=20)
    keep = dt.date(2026, 9, 30)
    r = admin.post(f'/api/programs/slow/{ta}', json={'keep_until': str(keep)})
    assert r.status_code == 200, r.text
    d = r.json()
    from app.programs import capacity
    with S() as db:
        cap = capacity(db, db.get(TeachingAssignment, ta))
        past_n = sum(1 for x, _ in cap['slots'] if x <= keep)
        old = [pl.topic for pl in db.query(PlanLesson).filter_by(assignment_id=ta).order_by(PlanLesson.seq)][:past_n]
    assert d['kept_past'] == past_n and [x['topic'] for x in d['preview'][:past_n]] == old       # keçmiş toxunulmur
    assert d['lessons'] == len(cap['slots'])                                                       # yuva = dərs
    assert d['practice'] > 0 and d['restore'] > 0 and d['bsq'] == 2
    assert d['section_review'] > 0 and d['spiral'] > 0 and d['ksq_prep'] == d['ksq'] - sum(1 for x in d['preview'][:past_n] if x['type'] == 'KSQ')
    for i, x in enumerate(d['preview']):                                            # hər yeni KSQ-dən əvvəl hazırlıq
        if i >= past_n and x['type'] == 'KSQ':
            assert d['preview'][i - 1]['topic'].startswith('KSQ-yə hazırlıq')
    types = [x['type'] for x in d['preview']]
    assert d['preview'][-1]['type'] == 'BSQ'
    ksq = [x['topic'] for x in d['preview'] if x['type'] == 'KSQ']
    assert all(t.startswith(f'KSQ-{i + 1}') or 'KSQ' in t for i, t in enumerate(ksq))
    r = admin.post(f'/api/programs/slow/{ta}', json={'keep_until': str(keep), 'apply': True})
    assert r.status_code == 200, r.text
    with S() as db:
        t = db.get(TeachingAssignment, ta)
        p1, p2 = db.get(PlanProgram, r.json()['program_id']), db.get(PlanProgram, r.json()['year2_program_id'])
        assert t.program_id == p1.id and p1.level == 'Zəif' and len(p2.data['lessons']) > 0
        n = db.query(PlanLesson).filter_by(assignment_id=ta).count()
        assert n == len(cap['slots'])


def test_plan_docx_export(world):
    """Perspektiv planın Word ixracı: başlıq, 8 sütunlu cədvəl, hər dərs bir sətir, məşq dərsi əsas mövzunun standartını alır."""
    import io
    from docx import Document
    as_, S = world
    admin, ta, cid, st, _ = setup(world)
    _plan(S, ta, n1=24, n2=20)
    admin.post(f'/api/programs/slow/{ta}', json={'keep_until': '2026-09-30', 'apply': True})
    r = admin.get(f'/api/plan/{ta}/docx')
    assert r.status_code == 200 and r.headers['content-type'].startswith('application/vnd.openxmlformats')
    d = Document(io.BytesIO(r.content))
    assert 'perspektiv plan' in d.paragraphs[1].text
    main = d.tables[1]
    assert [c.text for c in main.rows[0].cells][2] == 'Mövzu'
    with S() as db:
        n = db.query(PlanLesson).filter_by(assignment_id=ta).count()
    lesson_rows = [r_ for r_ in main.rows[1:] if r_.cells[0].text.strip().isdigit()]
    assert len(lesson_rows) == n


def test_lesson_query_derived():
    """Zəif tempin törəmə dərsləri: məşq/bərpa → əsas mövzunun testi, qarışıq təkrar → link yoxdur (docs/zeif-plan-test-linkleri-promtu.md)."""
    from app.bank.match import derived, lesson_query
    assert lesson_query('Məşq və möhkəmləndirmə: Kəsrlər', 'II BÖLMƏ – Kəsrlər') == ('Kəsrlər', 'II BÖLMƏ – Kəsrlər')
    assert lesson_query('Təkrar və bərpa: Faiz', 'Bazanın bərpası') == ('Faiz', None)
    for t in ('Bölmə təkrarı: Kəsrlər', 'Spiral təkrar: əvvəlki bölmələr', 'KSQ-yə hazırlıq: Kəsrlər',
              'Təkrar: I yarımil üzrə yekunlaşdırıcı tapşırıqlar (zəif qrup)'):
        assert lesson_query(t, 'X') is None and derived(t)
    assert lesson_query('Kəsrlər', 'II') == ('Kəsrlər', 'II') and not derived('Kəsrlər')


def test_apply_drops_stale_bank_link(world):
    """Proqram tətbiq olunanda mövzusu dəyişən dərsin köhnə test faylı seçimi silinir, dəyişməyənin seçimi qalır."""
    from app.models import LessonBankLink
    from app.programs import apply
    as_, S = world
    admin, ta, cid, st, _ = setup(world)
    _plan(S, ta, n1=6, n2=0)
    with S() as db:
        pls = db.query(PlanLesson).filter_by(assignment_id=ta).order_by(PlanLesson.seq).all()
        for pl in pls[:2]:
            db.add(LessonBankLink(plan_lesson_id=pl.id, file_id=None))
        lessons = [{'semester': 1, 'section': pl.section, 'topic': pl.topic if pl.seq == 1 else 'Yeni mövzu', 'standards': [],
                    'integration': None, 'resources': None, 'assessment': None, 'assessment_type': 'formativ', 'exam_no': None,
                    'tt_pages': None, 'tasks': []} for pl in pls]
        p = PlanProgram(school_id=1, owner_id=None, subject='Riyaziyyat', kind='fixed', weekly_hours=5, title='t',
                        data={'lessons': lessons})
        db.add(p)
        db.flush()
        user = db.get(__import__('app.models', fromlist=['User']).User, 1)
        apply(db, p, db.get(TeachingAssignment, ta), user)
        db.commit()
        kept = {k.plan_lesson_id for k in db.query(LessonBankLink)}
        assert kept == {pls[0].id}
