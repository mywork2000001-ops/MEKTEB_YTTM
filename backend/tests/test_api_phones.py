"""Şagird və valideyn telefonları: format, toplu daxil etmə, Excel/CSV idxalı, məxfilik (yalnız rəhbər/admin)."""
import io

import openpyxl
import pytest

from app.domain.phones import name_key, norm_phone


@pytest.mark.parametrize('raw, out', [
    ('050 123 45 67', '+994 50 123 45 67'), ('50-123-45-67', '+994 50 123 45 67'),
    ('994501234567', '+994 50 123 45 67'), ('+994 (55) 765 43 21', '+994 55 765 43 21'),
    ('501234567.0', '+994 50 123 45 67'), ('012 555 12 34', '+994 12 555 12 34'),
    ('+7 912 345 67 89', '+79123456789'), ('', None), (None, None)])
def test_norm_phone(raw, out):
    assert norm_phone(raw) == out


@pytest.mark.parametrize('raw', ['abc', '12345', '0991234', '+994 40 123 45 67'])
def test_norm_phone_bad(raw):
    with pytest.raises(ValueError):
        norm_phone(raw)


def test_name_key():
    assert name_key('Əliyev Rəşad Elçin oğlu') == name_key('əliyev  rəşad elçin')
    assert name_key('Quliyeva Aynur İlqar qızı') == name_key('Quliyeva Aynur Ilqar')


def _setup(world):
    as_, _ = world
    adm, t = as_('admin'), as_('ilqar')
    cid = adm.post('/api/classes', json={'name': 'X e'}).json()['id']
    a = adm.post('/api/students', json={'full_name': 'Əliyev Rəşad Elçin oğlu', 'class_id': cid}).json()['id']
    b = adm.post('/api/students', json={'full_name': 'Quliyeva Aynur İlqar qızı', 'class_id': cid}).json()['id']
    return adm, t, cid, a, b


def test_bulk_phones_and_privacy(world):
    adm, t, cid, a, b = _setup(world)
    assert t.get(f'/api/homeroom/{cid}/phones').status_code == 403                 # rəhbər deyil
    body = [{'student_id': a, 'phone': '050 111 22 33',
             'guardians': [{'name': 'Əliyeva Leyla', 'relation': 'ana', 'phone': '0557654321'}]},
            {'student_id': b, 'phone': None, 'guardians': []}]
    assert adm.put(f'/api/homeroom/{cid}/phones', json=body).status_code == 200
    r = adm.get(f'/api/homeroom/{cid}/phones').json()
    row = next(x for x in r['students'] if x['student_id'] == a)
    assert row['phone'] == '+994 50 111 22 33' and row['guardians'][0]['phone'] == '+994 55 765 43 21'
    assert r['filled'] == 1
    assert adm.put(f'/api/homeroom/{cid}/phones', json=[{'student_id': a, 'phone': 'xyz'}]).status_code == 422
    assert adm.put(f'/api/homeroom/{cid}/phones', json=[{'student_id': 99999}]).status_code == 404
    # fənn müəllimi şagird siyahısında telefonu görmür
    assert 'phone' not in adm.get('/api/students', params={'class_id': cid}).json()[0]


def _xlsx(rows):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_xlsx_dry_run_then_write(world):
    adm, _, cid, a, b = _setup(world)
    data = _xlsx([['Sinif X e – telefonlar'],
                  ['Şagird', 'Şagirdin telefonu', 'Ana', 'Ana telefonu', 'Ata', 'Ata telefonu'],
                  ['Əliyev Rəşad Elçin', '0501112233', 'Əliyeva Leyla', '055 765 43 21', '', ''],
                  ['Quliyeva Aynur', '', '', '', 'Quliyev İlqar', 'səhv'],
                  ['Naməlum Şagird', '0501234567', '', '', '', '']])
    files = {'file': ('tel.xlsx', data, 'application/octet-stream')}
    r = adm.post(f'/api/homeroom/{cid}/phones/import', files=files).json()
    assert r['dry_run'] and r['rows'] == 3 and r['matched'] == 1 and r['errors'] == 2
    assert adm.get(f'/api/homeroom/{cid}/phones').json()['filled'] == 0                # yoxlama yazmır
    files = {'file': ('tel.xlsx', data, 'application/octet-stream')}
    r = adm.post(f'/api/homeroom/{cid}/phones/import', params={'dry_run': False}, files=files).json()
    assert not r['dry_run']
    rows = {x['student_id']: x for x in adm.get(f'/api/homeroom/{cid}/phones').json()['students']}
    assert rows[a]['phone'] == '+994 50 111 22 33'
    assert rows[a]['guardians'] == [{'name': 'Əliyeva Leyla', 'relation': 'ana', 'phone': '+994 55 765 43 21'}]
    assert rows[b]['guardians'] == []                                                   # səhv nömrə – yazılmadı


def test_import_csv_and_template(world):
    adm, _, cid, a, _ = _setup(world)
    csv = 'Şagird;Ata;Ata telefonu\nƏliyev Rəşad Elçin oğlu;Əliyev Elçin;+994 70 222 33 44\n'.encode('utf-8-sig')
    r = adm.post(f'/api/homeroom/{cid}/phones/import', params={'dry_run': False},
                 files={'file': ('t.csv', csv, 'text/csv')}).json()
    assert r['matched'] == 1
    t = adm.get(f'/api/homeroom/{cid}/phones/template')
    assert t.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(t.content)).active
    vals = [[c.value for c in row] for row in ws.iter_rows()]
    assert vals[0][0] == 'Şagird' and any(v[0] == 'Əliyev Rəşad Elçin oğlu' and v[5] == '+994 70 222 33 44' for v in vals)
    bad = adm.post(f'/api/homeroom/{cid}/phones/import', files={'file': ('t.txt', b'x', 'text/plain')})
    assert bad.status_code == 400
