"""Fərdi hazırlıq: görüş (90 dəq. = 2 saat) üzrə perspektiv planı tətbiqdən kənarda hesablayır (docs/ferdi-perspektiv-plan-promtu.md §6).
Tətbiqin eyni funksiyalarını işlədir (meeting_parts, lesson_slots, expand_course, lessons_for, budget_summary) – ayrıca alqoritm yoxdur.

İstifadə (backend qovluğunda):
  python -m scripts.ferdi_plan --program buraxilis9 --from 2026-10-20 --to 2027-05-31 --time "1 17:00-18:30" --time "5 10:00-11:30" \\
        [--holidays tetil.txt] [--sections "V sinif*,VI sinif*"] [--mock-every 4] [--out plan.csv|plan.json|plan.docx]
  python -m scripts.ferdi_plan --class-id 123 --dry-run [--out ...]     # DB-dən oxuyur, YAZMIR
  python -m scripts.ferdi_plan --all-private --report                   # bütün fərdi qruplar: görüş/saat büdcəsi

--time "GÜN HH:MM-HH:MM": gün 0 (B.e.) … 6 (B.). --holidays: hər sətirdə YYYY-MM-DD. --mock-every – GÖRÜŞLƏRLƏ (saat = görüş × orta ölçü).
Heç bir rejim DB-yə yazmır."""
import argparse
import csv
import datetime as dt
import fnmatch
import json
import sys
from types import SimpleNamespace

DAYS = ['B.e.', 'Ç.a.', 'Ç.', 'C.a.', 'C.', 'Ş.', 'B.']


def parse_time(s: str) -> tuple[int, str, str]:
    day, rng = s.split()
    a, b = rng.split('-')
    return int(day), a, b


def load_holidays(path: str | None) -> dict[dt.date, str]:
    if not path:
        return {}
    out = {}
    for line in open(path, encoding='utf-8'):
        line = line.strip()
        if line and not line.startswith('#'):
            out[dt.date.fromisoformat(line[:10])] = line[10:].strip() or 'Tətil'
    return out


def builtin_template(key: str) -> dict:
    import json as _j

    from app.programs import DATA, _prep_template
    src = _j.loads((DATA / 'sinif_testleri.json').read_text(encoding='utf-8'))
    return _prep_template(src, {'buraxilis9': 9, 'buraxilis11': 11}[key])


def filter_sections(tpl: dict, patterns: str | None) -> dict:
    if not patterns:
        return tpl
    pats = [p.strip() for p in patterns.split(',') if p.strip()]
    tpl = dict(tpl)
    tpl['sections'] = [s for s in tpl['sections'] if any(fnmatch.fnmatch(s['section'], p) for p in pats)]
    return tpl


def times_map(times: list[tuple[int, str, str]]) -> dict[str, str]:
    """Görüşləri {"gün:sıra": alt vaxt} formatına açır (tətbiqdəki meeting_parts ilə; sıralar gündə 0-dan)."""
    from app.api.classes import meeting_parts
    tmap: dict[str, str] = {}
    for wd in sorted({t[0] for t in times}):
        p = 0
        for _, s, e in sorted(t for t in times if t[0] == wd):
            for s2, e2 in meeting_parts(s, e):
                tmap[f'{wd}:{p}'] = f'{s2}–{e2}'
                p += 1
    return tmap


def build(tpl: dict, tmap: dict[str, str], start: dt.date, end: dt.date, off: dict, sem1_end: dt.date | None = None):
    """(dərslər, yuvalar, hesabat) – yalnız verilənlərdən (DB-siz)."""
    from app.domain.calendar import lesson_slots
    from app.programs import budget_summary, expand_course, meeting_sizes
    slots_by_day: dict[int, list[int]] = {}
    for k in tmap:
        wd, p = map(int, k.split(':'))
        slots_by_day.setdefault(wd, []).append(p)
    slots = lesson_slots(slots_by_day, off, start, end)
    sizes = meeting_sizes(SimpleNamespace(times=tmap), slots)
    sem1_end = sem1_end or dt.date(start.year if start.month >= 9 else start.year - 1, 12, 27)
    lessons, warnings = expand_course(tpl, [d for d, _ in slots], sem1_end, sizes)
    n_top = sum(len(s['topics']) for s in tpl['sections'])
    rep = {'warnings': warnings, 'budget': budget_summary(slots, sizes, n_top * (min(sizes) if sizes else 1))}
    return lessons, slots, rep


def rows_of(lessons: list[dict], slots: list, tmap: dict[str, str]) -> list[dict]:
    rows, hour_in_day = [], 0
    for i, l in enumerate(lessons):
        d, p = slots[i] if i < len(slots) else (None, None)
        if d and i and slots[i - 1][0] == d:
            hour_in_day += 1
        else:
            hour_in_day = 1
        rows.append({'no': i + 1, 'date': d.isoformat() if d else '', 'day': DAYS[d.weekday()] if d else '',
                     'time': tmap.get(f'{d.weekday()}:{p}', '') if d else '', 'hour': hour_in_day,
                     'section': l.get('section') or '', 'topic': l['topic'], 'resource': l.get('resources') or '',
                     'assessment': l.get('assessment') or ''})
    return rows


COLS = [('no', '№'), ('date', 'Tarix'), ('day', 'Gün'), ('time', 'Vaxt'), ('hour', 'Saat'), ('section', 'Bölmə'),
        ('topic', 'Mövzu'), ('resource', 'Resurs'), ('assessment', 'Qiymətləndirmə')]


def summary_text(b: dict) -> str:
    s = f'Kurs: {b["from"]} – {b["to"]} · {b["meetings"]} görüş · {b["hours"]} saat'
    if b['need'] is not None:
        s += (f' · plan tələb edir {b["need"]} saat · ehtiyat {b["spare"]} saat' if b['spare'] >= 0
              else f' · plan tələb edir {b["need"]} saat · çatmır: {-b["spare"]} saat')
    return s


def write_out(path: str, rows: list[dict], rep: dict, title: str) -> None:
    b = rep['budget']
    if path.endswith('.json'):
        json.dump({'title': title, 'budget': b, 'warnings': rep['warnings'], 'rows': rows}, open(path, 'w', encoding='utf-8'),
                  ensure_ascii=False, indent=1)
    elif path.endswith('.csv'):
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow([summary_text(b)])
            w.writerow([h for _, h in COLS])
            for r in rows:
                w.writerow([r[k] for k, _ in COLS])
    elif path.endswith('.docx'):
        import docx
        doc = docx.Document()
        doc.add_heading(f'Perspektiv plan – {title}', 1)
        doc.add_paragraph(summary_text(b))
        for w_ in rep['warnings']:
            doc.add_paragraph(w_)
        t = doc.add_table(rows=1, cols=len(COLS))
        t.style = 'Table Grid'
        for c, (_, h) in zip(t.rows[0].cells, COLS):
            c.text = h
        for r in rows:
            for c, (k, _) in zip(t.add_row().cells, COLS):
                c.text = str(r[k])
        doc.save(path)
    else:
        raise SystemExit('--out: .csv, .json və ya .docx')


def db_plan(db, ta):
    """DB-dən: qrupun vaxtı, bayramlar, proqram → (dərslər, yuvalar, hesabat, vaxt xəritəsi). Yazmır."""
    from app.models import PlanProgram
    from app.programs import capacity, lessons_for
    p = db.get(PlanProgram, ta.program_id) if ta.program_id else None
    if not p:
        raise SystemExit(f'ta {ta.id}: proqram seçilməyib')
    lessons, rep = lessons_for(db, p, ta, ta.program_sections)
    return lessons, capacity(db, ta)['slots'], rep, dict(ta.times or {}), p.title


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--program', help='buraxilis9 | buraxilis11')
    ap.add_argument('--from', dest='start', type=dt.date.fromisoformat)
    ap.add_argument('--to', dest='end', type=dt.date.fromisoformat)
    ap.add_argument('--time', action='append', default=[], help='"GÜN HH:MM-HH:MM" (təkrar oluna bilər)')
    ap.add_argument('--holidays')
    ap.add_argument('--sections')
    ap.add_argument('--mock-every', type=int, help='aralıq sınaq – hər N görüşdən bir')
    ap.add_argument('--out')
    ap.add_argument('--class-id', type=int)
    ap.add_argument('--all-private', action='store_true')
    ap.add_argument('--report', action='store_true')
    ap.add_argument('--dry-run', action='store_true', help='DB-dən oxu, yazma (həmişə belədir)')
    a = ap.parse_args(argv)

    if a.program:
        if not (a.start and a.end and a.time):
            ap.error('--program ilə --from, --to və ən azı bir --time lazımdır')
        tpl = filter_sections(builtin_template(a.program), a.sections)
        times = [parse_time(t) for t in a.time]
        tmap = times_map(times)
        if a.mock_every:
            sizes_avg = len(tmap) / len(times)
            tpl['mock_every'] = max(1, round(a.mock_every * sizes_avg))
        lessons, slots, rep = build(tpl, tmap, a.start, a.end, load_holidays(a.holidays))
        rows, title = rows_of(lessons, slots, tmap), a.program
    elif a.class_id or a.all_private:
        from sqlalchemy import select

        from app.db import SessionLocal
        from app.models import School, SchoolClass, TeachingAssignment
        db = SessionLocal()
        try:
            q = select(TeachingAssignment).join(SchoolClass, SchoolClass.id == TeachingAssignment.class_id).where(
                TeachingAssignment.archived_at.is_(None), TeachingAssignment.times.is_not(None))
            if a.class_id:
                q = q.where(TeachingAssignment.class_id == a.class_id)
            else:
                q = q.join(School, School.id == SchoolClass.school_id).where(School.kind == 'private')
            tas = list(db.scalars(q))
            if not tas:
                raise SystemExit('Uyğun qrup tapılmadı')
            for ta in tas:
                c = db.get(SchoolClass, ta.class_id)
                try:
                    lessons, slots, rep, tmap, title = db_plan(db, ta)
                except SystemExit as e:
                    print(f'{c.name}: {e}')
                    continue
                print(f'{c.name} / {ta.subject}: {summary_text(rep["budget"])}')
                if a.class_id:
                    rows = rows_of(lessons, slots, tmap)
                    title = f'{c.name} – {title}'
        finally:
            db.close()
        if a.all_private:
            return 0
    else:
        ap.error('--program, --class-id və ya --all-private seçin')
        return 2
    print(summary_text(rep['budget']))
    for w in rep['warnings']:
        print('!', w)
    if a.out:
        write_out(a.out, rows, rep, title)
        print('Yazıldı:', a.out)
    else:
        for r in rows[:12]:
            print(r['no'], r['date'], r['day'], r['time'], r['topic'])
    return 0


if __name__ == '__main__':
    sys.exit(main())
