"""Sinif matrisinin («kim nəyi işləyib») rəsmi çap formaları – A4 albom: HTML (PDF üçün – api/printing.render_pdf) və Excel (openpyxl).
Rəsmi sənəd yalnız Azərbaycan dilində. Nəşr olunmamış skan balı mötərizədə."""
from __future__ import annotations

import io
from html import escape as _e

DOC = 'Kim nəyi işləyib – sinif üzrə test və sınaq nəticələri'
KIND_LABEL = {'movzu': 'Mövzu testləri', 'sinaq': 'Sınaqlar', 'all': 'Mövzu testləri, sınaqlar, KSQ/BSQ'}


def _d(s: str | None) -> str:
    return '.'.join(s[:10].split('-')[::-1]) if s else '—'


def _n(v, d=0) -> str:
    return '—' if v is None else f'{v:.{d}f}'.replace('.', ',')


def cell_text(x: dict) -> str:
    st = x['status']
    if st == 'yazıb':
        return _n(x['pct'])
    if st == 'nəşr olunmayıb':
        return f"({_n(x['pct'])})" if x['pct'] is not None else 'nəşr olunmayıb'
    return {'yazmayıb': 'yazmayıb', 'təyin olunmayıb': '—', 'gözlənilir': 'gözlənilir', 'işlənməyib': 'işlənməyib'}.get(st, '—')


def col_head(c: dict) -> str:
    return ((c['no'] + ' ') if c.get('no') else '') + (c['title'] or '')


def period_text(m: dict) -> str:
    p = m['period']
    return f"{_d(p['from'])} – {_d(p['to'])}" if p['from'] or p['to'] else 'bütün dövr'


def matrix_html(m: dict, school: str, teacher: str) -> str:
    cols = m['columns']
    head = ''.join(f"<th>{_e(col_head(c))}<br><small>{_e(_d(c['date']))}</small></th>" for c in cols)
    body = []
    for n, r in enumerate(m['rows'], 1):
        tds = ''.join(f"<td class=\"c\">{_e(cell_text(r['cells'][c['key']]))}</td>" for c in cols)
        body.append(f"<tr><td>{n}</td><td class=\"nm\">{_e(r['full_name'])}</td>{tds}"
                    f"<td class=\"c\">{r['wrote']}/{r['given']}</td><td class=\"c\">{_e(_n(r['avg_pct']))}</td></tr>")
    foot1 = ''.join(f"<td class=\"c\">{c['wrote_count']}/{c['assigned_count']}</td>" for c in cols)
    foot2 = ''.join(f"<td class=\"c\">{_e(_n(c['avg_pct']))}</td>" for c in cols)
    return (
        '<!doctype html><html lang="az"><head><meta charset="utf-8"><title>' + _e(DOC) + '</title><style>'
        '@page{size:A4 landscape;margin:10mm}body{font:10px/1.3 "Times New Roman","Liberation Serif",serif;color:#000}'
        'h1{font-size:14px;text-align:center;margin:0 0 2px}p{margin:2px 0;text-align:center}'
        'table{border-collapse:collapse;width:100%;margin-top:6px}th,td{border:1px solid #000;padding:2px 3px}'
        'th{background:#D9D9D9}thead{display:table-header-group}tr{page-break-inside:avoid}.c{text-align:center}'
        '.nm{white-space:nowrap}.sign{margin-top:20px}small{font-weight:normal}</style></head><body>'
        f'<p><b>{_e(school)}</b></p><h1>{_e(DOC)}</h1>'
        f'<p>{_e(m["class_name"])} · {_e(m["subject"])} · {_e(KIND_LABEL[m["kind"]])} · {_e(period_text(m))}'
        f' · müəllim: {_e(teacher)}</p>'
        f'<table><thead><tr><th>№</th><th>Şagird</th>{head}<th>Yazıb/təyin</th><th>Orta, %</th></tr></thead><tbody>'
        + ''.join(body) +
        f'<tr><td></td><td><b>Yazan / təyin olunan</b></td>{foot1}<td></td><td></td></tr>'
        f'<tr><td></td><td><b>Orta, %</b></td>{foot2}<td></td><td></td></tr></tbody></table>'
        '<p class="c" style="text-align:left">(xx) – nəticə şagirdə hələ nəşr olunmayıb.</p>'
        f'<p class="sign" style="text-align:left">Fənn müəllimi: {_e(teacher)} ______________ (imza)</p></body></html>')


def matrix_xlsx(m: dict, school: str, teacher: str, deputy: str = '') -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    cols = m['columns']
    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    grey = PatternFill('solid', fgColor='D9D9D9')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Matris'
    n = len(cols) + 4
    lines = [(school, True, 11), (DOC, True, 12),
             (f'{m["class_name"]} · {m["subject"]} · {KIND_LABEL[m["kind"]]} · {period_text(m)} · müəllim: {teacher}', False, 10)]
    for i, (txt, bold, size) in enumerate(lines, 1):
        ws.cell(i, 1, txt).font = Font(bold=bold, size=size)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=n)
        ws.cell(i, 1).alignment = Alignment(horizontal='center')
    hr = 5
    heads = ['№', 'Şagird'] + [f"{col_head(c)}\n{_d(c['date'])}" for c in cols] + ['Yazıb/təyin', 'Orta, %']
    for j, h in enumerate(heads, 1):
        c = ws.cell(hr, j, h)
        c.font, c.fill, c.border = Font(bold=True), grey, border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[hr].height = 48

    def put(row: list):
        ws.append(row)
        for c in ws[ws.max_row]:
            c.border = border
            if c.column != 2:
                c.alignment = Alignment(horizontal='center')

    for i, r in enumerate(m['rows'], 1):
        put([i, r['full_name']] + [cell_text(r['cells'][c['key']]) for c in cols] + [f"{r['wrote']}/{r['given']}", r['avg_pct']])
    put(['', 'Yazan / təyin olunan'] + [f"{c['wrote_count']}/{c['assigned_count']}" for c in cols] + ['', ''])
    put(['', 'Orta, %'] + [c['avg_pct'] for c in cols] + ['', ''])
    last = ws.max_row + 2
    ws.cell(last, 1, '(xx) – nəticə şagirdə hələ nəşr olunmayıb.').font = Font(italic=True, size=9)
    last += 2
    for role, name in (('Fənn müəllimi', teacher), ('Direktor müavini (tədris işləri üzrə)', deputy)):
        ws.cell(last, 1, f'{role}:')
        ws.cell(last, 3, '______________ (imza)')
        ws.cell(last, 5, name or '______________')
        last += 2
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 32
    for i in range(3, n + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{hr}:{hr}'
    ws.freeze_panes = ws.cell(hr + 1, 3)
    ws.oddFooter.right.text = 'Səhifə &P / &N'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
