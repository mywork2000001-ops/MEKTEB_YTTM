"""Şagird siyahısının importu. İki format:
1) «Buraxılış nəticələri - riyaziyyat (Həsənov F.).xlsx» – vərəq başına bir sinif, «№ | Soyadı, adı, ata adı | …»;
2) DİM ixracı «TOM-şagirdlər buraxılış balları.xlsx» – bir vərəqdə bütün siniflər, «Tədris sinfi» sütunu ilə
   (başlıqlar «Tədrİs sİnfİ» kimi nöqtəli İ ilə yazılıb, «Tədris dili» iki dəfə: dil adı və bal).
Pinkod, Uşaq İD, müəssisə İD, rayon və məktəb sütunları heç vaxt oxunmur və saxlanmır."""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass
from pathlib import Path

FORBIDDEN = ('pinkod', 'pin kod', 'uşaq id', 'müəssisə', 'rayon')

COLS = {                      # normallaşdırılmış başlığın əvvəli -> sahə
    'soyadı': 'name',
    'cins': 'gender',
    'doğum tarixi': 'birth_date',
    'tədris sinfi': 'class_name',
    'tədris dili': 'score_language',   # dil adı olan sütun da belə adlanır – rəqəmli olan götürülür
    'riyaziyyat': 'score_math',
    'xarici dil': 'score_foreign',
}
NUMERIC = {'score_language', 'score_math', 'score_foreign'}
GENDER = {'qadın': 'Qız', 'kişi': 'Oğlan', 'qız': 'Qız', 'oğlan': 'Oğlan'}


def norm_head(h) -> str:
    """«Tədrİs sİnfİ», «Pİnkod» -> «tədris sinfi», «pinkod» (İ böyük/kiçik və birləşən nöqtə aradan qalxır)."""
    return ' '.join(str(h or '').replace('İ', 'i').replace('̇', '').lower().split())


def az_title(s: str) -> str:
    """«ƏHMƏDZADƏ GÜNAY MAARİF» -> «Əhmədzadə Günay Maarif» (I -> ı, İ -> i qaydası ilə)."""
    def w(x):
        low = x.replace('̇', '').replace('İ', 'i').replace('I', 'ı').lower()
        return (low[0].replace('i', 'İ').replace('ı', 'I').upper() if low[0] in 'iı' else low[0].upper()) + low[1:]
    return ' '.join(w(x) for x in str(s).split())


@dataclass
class ImportedStudent:
    no: int
    name: str
    gender: str | None
    birth_date: dt.date | None
    score_language: float | None
    score_math: float | None
    score_foreign: float | None
    class_name: str | None = None      # yalnız DİM formatında: «10 b»

    @property
    def score_total(self) -> float | None:
        v = [self.score_language, self.score_math, self.score_foreign]
        return round(sum(v), 1) if all(x is not None for x in v) else None


def _date(v) -> dt.date | None:
    if v in (None, ''):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    for fmt in ('%d/%m/%Y', '%d.%m.%Y', '%Y-%m-%d'):
        try:
            return dt.datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f'Doğum tarixi oxunmadı: {v!r}')


def _num(v) -> float | None:
    if v in (None, ''):
        return None
    return round(float(str(v).replace(',', '.')), 1)


def read_sheet(ws) -> tuple[list[ImportedStudent], list[str]]:
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows) if r and any(norm_head(v).startswith('soyadı') for v in r)), None)
    if hi is None:
        raise ValueError(f'«{ws.title}»: başlıq sətri tapılmadı')
    head = [norm_head(v) for v in rows[hi]]
    body = [r for r in rows[hi + 1:] if r and str(r[0] or '').strip().isdigit()]
    idx = {}
    for key, f in COLS.items():
        cand = [i for i, h in enumerate(head) if h.startswith(key)]
        if f in NUMERIC:     # eyni adlı sütunlardan rəqəmli olanı
            cand = [i for i in cand if any(isinstance(r[i], (int, float)) for r in body[:5])] or cand
        if cand:
            idx[f] = cand[-1]
    if 'name' not in idx:
        raise ValueError(f'«{ws.title}»: «Soyadı, adı, ata adı» sütunu yoxdur')
    out, warnings = [], []
    skipped = [rows[hi][i] for i, h in enumerate(head) if any(k in h for k in FORBIDDEN)]
    if skipped:
        warnings.append(f'«{ws.title}»: məxfi sütunlar oxunmadı: {", ".join(map(str, skipped))}')
    for r in body:
        g = lambda f: r[idx[f]] if f in idx else None
        if not g('name'):
            continue
        try:
            name = str(g('name'))
            out.append(ImportedStudent(
                int(str(r[0]).strip()), az_title(name) if name.isupper() else ' '.join(name.split()),
                GENDER.get(norm_head(g('gender')), g('gender') or None), _date(g('birth_date')),
                _num(g('score_language')), _num(g('score_math')), _num(g('score_foreign')),
                str(g('class_name')).strip() if g('class_name') else None))
        except ValueError as e:
            warnings.append(f'«{ws.title}» №{r[0]}: {e}')
    return out, warnings


def read_workbook(path: str | Path, sheets: list[str] | None = None) -> tuple[dict[str, list[ImportedStudent]], list[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    result, warnings = {}, []
    for name in sheets or wb.sheetnames:
        if name not in wb.sheetnames:
            warnings.append(f'Vərəq tapılmadı: «{name}»')
            continue
        try:
            result[name], w = read_sheet(wb[name])
            warnings += w
        except ValueError as e:
            warnings.append(str(e))
    return result, warnings
