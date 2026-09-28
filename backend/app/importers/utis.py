"""Əsas şagird siyahısı: UTİS ixracı («Utis_siyahi (27).xlsx»).
Sütunlar: NO | Tədris sinfi | Soyadı | Adı | Atasının adı | Uşaq İD | Cinsi | Doğum tarixi | ŞV seriya… | Seriya/nömrə | Pinkod.
Buraxılış balları DİM ixracından («TOM-şagirdlər buraxılış balları.xlsx») Uşaq İD üzrə YALNIZ yaddaşda birləşdirilir.
Uşaq İD, şəxsiyyət vəsiqəsi və pinkod nəticəyə düşmür – heç yerdə saxlanmır."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from .students import ImportedStudent, norm_head, _date, GENDER

NEED = {'tədris sinfi': 'cls', 'soyadı': 'last', 'adı': 'first', 'atasının adı': 'father',
        'cinsi': 'gender', 'doğum tarixi': 'birth', 'uşaq id': 'kid'}


@dataclass
class UtisRoster:
    classes: dict[str, list[ImportedStudent]]      # «10 b» -> şagirdlər (UTİS-dəki ardıcıllıqla)
    without_scores: list[tuple[str, str]]           # (sinif, ad) – DİM faylında balı yoxdur
    warnings: list[str]


def _rows(path, sheet=None):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active
    return list(ws.iter_rows(values_only=True))


def _dim_scores(path) -> dict[int, tuple]:
    """Uşaq İD -> (dil, riyaziyyat, xarici dil). Yalnız bu funksiyanın daxilində və qaytarılan lüğətdə yaşayır."""
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = {}
    for ws in wb:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head = [norm_head(h) for h in rows[0]]
        if 'uşaq id' not in head or 'riyaziyyat' not in head:
            continue
        kid = head.index('uşaq id')
        num = lambda key: [i for i, h in enumerate(head) if h == key][-1]
        lang, math, foreign = num('tədris dili'), num('riyaziyyat'), num('xarici dil')
        for r in rows[1:]:
            if r and r[kid]:
                out[int(r[kid])] = tuple(round(float(r[i]), 1) if r[i] not in (None, '') else None
                                         for i in (lang, math, foreign))
    return out


def class_label(utis_cls: str) -> str:
    """UTİS «10 a 1» / «11 b» -> interfeys «X a» / «XI b» (sonrakı rəqəm UTİS-in daxili bölgüsüdür)."""
    parts = utis_cls.split()
    roman = {'10': 'X', '11': 'XI', '9': 'IX'}.get(parts[0], parts[0])
    return f'{roman} {parts[1]}' if len(parts) > 1 else roman


def read_utis(path: str | Path, dim_path: str | Path | None = None) -> UtisRoster:
    rows = _rows(path)
    hi = next(i for i, r in enumerate(rows) if r and any(norm_head(v) == 'soyadı' for v in r))
    head = [norm_head(v) for v in rows[hi]]
    ix = {f: head.index(k) for k, f in NEED.items() if k in head}
    missing = [k for k, f in NEED.items() if f not in ix and f != 'kid']
    if missing:
        raise ValueError('UTİS faylında sütun yoxdur: ' + ', '.join(missing))
    scores = _dim_scores(dim_path) if dim_path else {}
    classes, without, warnings = {}, [], []
    for r in rows[hi + 1:]:
        if not r or not r[ix['last']]:
            continue
        name = ' '.join(f'{r[ix["last"]]} {r[ix["first"]]} {r[ix["father"]] or ""}'.split())
        cls = ' '.join(str(r[ix['cls']]).split())
        kid = r[ix['kid']] if 'kid' in ix else None
        sc = scores.get(int(kid)) if kid and scores else None
        if dim_path and sc is None:
            without.append((cls, name))
        try:
            birth = _date(r[ix['birth']])
        except ValueError as e:
            warnings.append(f'{name}: {e}')
            birth = None
        lst = classes.setdefault(cls, [])
        lst.append(ImportedStudent(len(lst) + 1, name, GENDER.get(norm_head(r[ix['gender']]), r[ix['gender']]),
                                   birth, *(sc or (None, None, None)), class_name=cls))
    return UtisRoster(classes, without, warnings)
