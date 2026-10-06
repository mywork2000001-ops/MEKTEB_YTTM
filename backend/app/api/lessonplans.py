"""Gündəlik planlaşdırma (ARTİ): perspektiv plan üzrə hər dərs yuvası üçün süni intellektlə gündəlik plan.

Müəllim öz API açarını seçir (Gemini, OpenRouter, OpenAI, Anthropic, Groq, DeepSeek, OpenAI-uyğun).
Gün/həftə görünüşü, hazırlama (bir-bir – həftə üçün interfeys ardıcıl çağırır), redaktə, silmə, Word (.docx)."""
from __future__ import annotations

import datetime as dt
import io
import re

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import ai
from ..services import ws_cond
from ..db import get_db
from ..deps import staff
from ..domain import daily_plan as dp
from ..domain.plan import slot_at, view_range
from ..models import DailyPlan, JournalEntry, User
from ..security import secret_decrypt, secret_encrypt
from ..services import journal_entries, own_assignment, plan_ctx, roster, taught_lesson
from .common import audit, settings_unlocked
from .plan import WEEKDAYS, bell

router = APIRouter(prefix='/api', tags=['daily-plans'])
SCHOOL_DEFAULT = 'Tərtər şəhər Rafiq Nuriyev adına 6 nömrəli tam orta ümumtəhsil məktəbi'


def header_school(user: User) -> str:
    """Başlıqda məktəbin adı: müəllimin seçimi, yoxdursa rəsmi ad (XI peşə də daxil – eyni məktəb)."""
    if (h := (user.ai_settings or {}).get('header_school')):
        return h
    from sqlalchemy.orm import object_session
    from ..models import School
    db = object_session(user)
    sc = db.get(School, user.school_id) if db is not None and user.school_id else None
    return sc.name if sc and sc.name else SCHOOL_DEFAULT
DAYS_FULL = ['Bazar ertəsi', 'Çərşənbə axşamı', 'Çərşənbə', 'Cümə axşamı', 'Cümə', 'Şənbə', 'Bazar']


# ---------------------------------------------------------------- müəllimin API açarı
def _ai_out(user: User) -> dict:
    s = user.ai_settings or {}
    key = secret_decrypt(s.get('key'))
    return {'provider': s.get('provider'), 'model': s.get('model'), 'base_url': s.get('base_url'),
            'has_key': bool(key), 'key_mask': ('••••' + key[-4:]) if key and len(key) > 8 else ('••••' if key else None),
            'providers': [{'id': k, 'label': v['label'], 'models': v['models'], 'key_url': v.get('key_url'),
                           'hint': v.get('hint'), 'custom': v['base'] is None} for k, v in ai.PROVIDERS.items()]}


def ai_config(user: User) -> dict:
    s = user.ai_settings or {}
    key = secret_decrypt(s.get('key'))
    if not s.get('provider') or not key or not s.get('model'):
        raise HTTPException(400, 'Süni intellekt qoşulmayıb: Tənzimləmələr → «Süni intellekt» bölməsində API açarınızı daxil edin')
    if s['provider'] not in ai.PROVIDERS:
        raise HTTPException(400, 'Provayder tanınmadı – tənzimləmələri yenidən saxlayın')
    return {'provider': s['provider'], 'model': s['model'], 'base_url': s.get('base_url'), 'api_key': key}


class AiIn(BaseModel):
    provider: str
    model: str = Field(min_length=2, max_length=120)
    api_key: str | None = Field(None, max_length=400)          # boş – köhnə açar qalır
    base_url: str | None = Field(None, max_length=300)


@router.get('/ai/settings')
def get_ai(user: User = Depends(staff)):
    return _ai_out(user)


@router.put('/ai/settings')
def put_ai(body: AiIn, user: User = Depends(settings_unlocked), db: Session = Depends(get_db)):
    if body.provider not in ai.PROVIDERS:
        raise HTTPException(400, 'Naməlum provayder')
    old = dict(user.ai_settings or {})
    base = None
    if ai.PROVIDERS[body.provider]['base'] is None:
        if not body.base_url:
            raise HTTPException(400, '«Başqa» provayder üçün ünvan (base URL) lazımdır')
        base = ai.check_base_url(body.base_url)
    key = (body.api_key or '').strip()
    if key and (len(key) < 12 or re.search(r'\s', key)):
        raise HTTPException(400, 'API açarı düzgün görünmür')
    if not key:
        if old.get('provider') != body.provider or not old.get('key'):
            raise HTTPException(400, 'API açarını daxil edin')
        enc = old['key']
    else:
        enc = secret_encrypt(key)
    user.ai_settings = {'provider': body.provider, 'model': body.model.strip(), 'base_url': base, 'key': enc,
                        'header_school': old.get('header_school')}
    audit(db, user, 'update', 'ai_settings', user.id, provider=body.provider, model=body.model)
    db.commit()
    return _ai_out(user)


@router.delete('/ai/settings')
def del_ai(user: User = Depends(settings_unlocked), db: Session = Depends(get_db)):
    hs = (user.ai_settings or {}).get('header_school')
    user.ai_settings = {'header_school': hs} if hs else None
    audit(db, user, 'delete', 'ai_settings', user.id)
    db.commit()
    return _ai_out(user)


@router.post('/ai/test')
def test_ai(user: User = Depends(staff)):
    cfg = ai_config(user)
    txt = ai.complete(cfg, 'Qısa cavab ver.', 'Yalnız bu sözü yaz: HAZIR', json_mode=False, max_tokens=400)
    return {'ok': True, 'reply': txt.strip()[:80], 'model': cfg['model']}


# ---------------------------------------------------------------- kontekst (perspektiv plan + jurnal)
def _slot(ctx, d: dt.date, period: int):
    s = slot_at(ctx.slots, d, period)
    if not s:
        raise HTTPException(404, 'Bu tarixdə və saatda sizin dərsiniz yoxdur')
    return s


def _composition(db: Session, ctx) -> dict | None:
    """Sinfin səviyyə tərkibi (müəllimin əl ilə təyini və ya nəticələr/IX balı) – diferensial yanaşma üçün. Adlar yoxdur."""
    try:
        from collections import Counter
        from ..analytics import analyze
        from .analytics import _period
        rows = analyze(db, ctx, _period(ctx, None, None, None))['students']
        cnt = Counter(r['level'] or 'Məlum deyil' for r in rows)
        return {k: cnt[k] for k in ('Güclü', 'Orta', 'Zəif', 'Məlum deyil') if cnt[k]}
    except Exception:                                                    # noqa: BLE001 – plan yenə hazırlansın
        return None


def _context(db: Session, user: User, ctx, s, notes: str | None) -> tuple[dict, object]:
    entry = db.scalar(select(JournalEntry).where(JournalEntry.assignment_id == ctx.ta.id, JournalEntry.date == s.date,
                                                 JournalEntry.period == s.period))
    pl = taught_lesson(ctx, s, entry)
    if pl is None:
        raise HTTPException(400, 'Bu dərs üçün perspektiv planda mövzu yoxdur (plan yüklənməyib və ya ilin sonuna sığmır)')
    i = next(k for k, x in enumerate(ctx.lessons) if x.id == pl.id)
    k = ctx.slots.index(s)
    prev_s = ctx.slots[k - 1] if k > 0 else None
    prev_pl = ctx.lessons[i - 1] if i > 0 else None
    nxt = ctx.lessons[i + 1] if i + 1 < len(ctx.lessons) else None
    prev_entry = db.scalar(select(JournalEntry).where(JournalEntry.assignment_id == ctx.ta.id, JournalEntry.homework.is_not(None),
                                                      (JournalEntry.date < s.date) | ((JournalEntry.date == s.date) & (JournalEntry.period < s.period)))
                           .order_by(JournalEntry.date.desc(), JournalEntry.period.desc()).limit(1))
    sec = [x for x in ctx.lessons if x.section == pl.section] if pl.section else []
    exam = None
    for j, x in enumerate(ctx.lessons[i + 1:], 1):
        if x.assessment_type in ('KSQ', 'BSQ'):
            exam = f"{x.assessment_type}{('-' + str(x.exam_no)) if x.exam_no else ''} – {j} dərs sonra"
            break
    hw_t = next((t for t in pl.tasks or [] if t.get('kind') == 'ev'), None)
    hw, hw_nums = None, []
    if hw_t:
        hw_nums = [str(hw_t.get('start'))] + ([str(hw_t['end'])] if hw_t.get('end') and hw_t.get('end') != hw_t.get('start') else [])
        hw = f"{pl.tt_pages + ', ' if pl.tt_pages else ''}№ {'–'.join(hw_nums)}"
    c = {'kind': ctx.cls.kind, 'levels': _composition(db, ctx), 'school': header_school(user), 'teacher': user.full_name, 'subject': ctx.ta.subject,
         'class_name': ctx.cls.name, 'group': ctx.cls.kind == 'qrup', 'students': len(roster(db, ctx.ta)),
         'date': s.date.isoformat(), 'date_text': s.date.strftime('%d.%m.%Y'), 'weekday': DAYS_FULL[s.date.weekday()],
         'period': s.period, 'time': bell(db, ctx.cls, s.period, ctx.ta, s.date), 'minutes': dp.LESSON_MIN,
         'semester': pl.semester, 'seq': pl.seq, 'total': len(ctx.lessons), 'section': pl.section,
         'section_pos': (sec.index(pl) + 1) if sec else None, 'section_len': len(sec) or None,
         'topic': pl.topic, 'standards': list(pl.standards or []), 'assessment_type': pl.assessment_type,
         'exam_no': pl.exam_no, 'assessment': pl.assessment, 'integration': pl.integration, 'resources': pl.resources,
         'tt_pages': pl.tt_pages, 'tasks': dp.tasks_text(pl.tasks),
         'continues_from': bool(prev_s and prev_s.held and prev_s.index == s.index),
         'continues_next': bool(s.held),
         'prev_topic': prev_pl.topic if prev_pl else None, 'prev_homework': prev_entry.homework if prev_entry else None,
         'next_topic': nxt.topic if nxt else None, 'next_exam': exam, 'notes': (notes or '').strip() or None,
         'homework_plan': hw, 'homework_nums': hw_nums, 'review_topics': _review_topics(db, ctx, i)}
    from .plan import attach_p0010
    x = {'id': pl.id, 'topic': pl.topic, 'section': pl.section, 'assessment_type': pl.assessment_type}
    attach_p0010(db, ctx.cls, [x])              # X–XI: mövzuya uyğun P0010 test faylı (yoxlama və ev tapşırığı üçün)
    c['p0010'] = (x.get('p0010') or {}).get('label')
    return c, pl


def _review_topics(db: Session, ctx, i: int) -> list[str]:
    """Müəllimin «təkrar» / «qismən» qeyd etdiyi əvvəlki mövzular (son 5) – dərsin əvvəlində qısa təkrar üçün."""
    from ..models import TopicProgress
    prev = {x.id: x for x in ctx.lessons[:i]}
    rows = db.scalars(select(TopicProgress).where(TopicProgress.assignment_id == ctx.ta.id,
                                                  TopicProgress.status.in_(('təkrar', 'qismən')),
                                                  TopicProgress.plan_lesson_id.in_(list(prev) or [0])))
    found = sorted((prev[m.plan_lesson_id].seq, prev[m.plan_lesson_id].topic, m.status) for m in rows)
    return [f"№{s} {t} ({'qismən keçilib' if st == 'qismən' else 'mənimsəmə zəifdir'})" for s, t, st in found[-5:]]


def _brief(p: DailyPlan | None) -> dict | None:
    return p and {'id': p.id, 'updated_at': p.updated_at, 'model': p.model, 'edited': p.edited, 'topic': p.topic}


def _full(p: DailyPlan, ctx) -> dict:
    return {**_brief(p), 'date': p.date, 'period': p.period, 'content': p.content, 'notes': p.notes,
            'provider': p.provider, 'class_name': ctx.cls.name, 'subject': ctx.ta.subject,
            'weekday': DAYS_FULL[p.date.weekday()], 'warnings': (p.content or {}).get('_warnings', []),
            'meta': {**(p.content or {}).get('_meta', {}), 'school': header_school(ctx.ta.teacher)}}


# ---------------------------------------------------------------- gündəlik planlar
META_KEYS = ('school', 'teacher', 'subject', 'class_name', 'date_text', 'weekday', 'period', 'time', 'minutes',
             'semester', 'seq', 'total', 'section', 'topic', 'assessment_type', 'exam_no', 'tt_pages', 'tasks', 'resources')


def _my_assignments(db: Session, user: User):
    from ..models import SchoolClass, TeachingAssignment
    return list(db.scalars(select(TeachingAssignment).join(SchoolClass).where(
        TeachingAssignment.teacher_id == user.id, ws_cond(user), TeachingAssignment.archived_at.is_(None),
        SchoolClass.archived_at.is_(None)).order_by(SchoolClass.name)))


def _range(view: str, d: dt.date) -> tuple[dt.date, dt.date]:
    if view == 'day':
        return d, d
    if view == 'week':
        a = d - dt.timedelta(days=d.weekday())
        return a, a + dt.timedelta(days=4)
    raise HTTPException(400, 'görünüş: day, week')


def _items(db: Session, tas, a: dt.date, b: dt.date) -> list[dict]:
    items = []
    for ta in tas:
        ctx = plan_ctx(db, ta)
        slots = [s for s in ctx.slots if a <= s.date <= b]
        if not slots:
            continue
        saved = {(p.date, p.period): p for p in db.scalars(select(DailyPlan).where(
            DailyPlan.assignment_id == ta.id, DailyPlan.date >= a, DailyPlan.date <= b))}
        entries = journal_entries(db, ta.id, a, b)
        for s in slots:
            pl = taught_lesson(ctx, s, entries.get((s.date, s.period)))
            items.append({'ta_id': ta.id, 'class_name': ctx.cls.name, 'subject': ta.subject, 'date': s.date,
                          'weekday': WEEKDAYS[s.date.weekday()], 'period': s.period, 'time': bell(db, ctx.cls, s.period, ctx.ta, s.date),
                          'held': s.held,
                          'lesson': pl and {'seq': pl.seq, 'topic': pl.topic, 'section': pl.section,
                                            'standards': pl.standards, 'assessment_type': pl.assessment_type,
                                            'exam_no': pl.exam_no, 'tt_pages': pl.tt_pages,
                                            'tasks': dp.tasks_text(pl.tasks)},
                          'plan': _brief(saved.get((s.date, s.period)))})
    items.sort(key=lambda i: (i['date'], i['period'], i['class_name']))
    return items


@router.get('/daily-plans-list')
def list_all(ta: str = 'all', view: str = 'day', date: dt.date | None = None, user: User = Depends(staff),
             db: Session = Depends(get_db)):
    """Gün və ya dərs həftəsi; bütün siniflər (ta=all) və ya bir sinif/qrup (ta=<id>) – tarix və saat sırası ilə."""
    from ..services import today
    d = date or today()
    a, b = _range(view, d)
    tas = _my_assignments(db, user) if ta == 'all' else [own_assignment(db, user, int(ta) if ta.isdigit() else 0)]
    s = user.ai_settings or {}
    return {'from': a, 'to': b, 'weekday': DAYS_FULL[d.weekday()], 'items': _items(db, tas, a, b),
            'header_school': header_school(user),
            'ai': {'configured': bool(s.get('key')), 'provider': s.get('provider'), 'model': s.get('model')}}


@router.get('/daily-plans/{ta_id}')
def list_plans(ta_id: int, view: str = 'week', date: dt.date | None = None, user: User = Depends(staff),
               db: Session = Depends(get_db)):
    """Bir sinif üzrə (köhnə ünvan) – list_all ilə eyni."""
    return list_all(str(ta_id), view, date, user, db)


def _blank(db: Session, user: User, ctx, s) -> dict:
    """Süni intellektsiz şablon: perspektiv plandan doldurulur (mövzu, altstandart, meyar, ev tapşırığı), qalanı boş."""
    c, pl = _context(db, user, ctx, s, None)
    res = [x.strip() for x in re.split(r'[;\n]', pl.resources or '') if x.strip()]
    content = {'standartlar': [{'kod': k, 'metn': ''} for k in c['standards']], 'telim_neticeleri': [],
               'acar_anlayislar': [], 'inteqrasiya': c['integration'] or '', 'is_formalari': [], 'is_usullari': [],
               'resurslar': res, 'tedqiqat_suali': '', 'merheleler': [], 'diferensial': {'destek': '', 'inkisaf': ''},
               'qiymetlendirme': {'meyarlar': [pl.assessment] if pl.assessment else [], 'usul': '', 'vasite': '',
                                  'rubrika': [], 'spesifikasiya': []},
               'refleksiya': [], 'ev_tapsirigi': c['homework_plan'] or '', 'muellim_ucun_qeyd': '',
               '_meta': {k: c[k] for k in META_KEYS}, '_warnings': []}
    return {'id': None, 'blank': True, 'topic': pl.topic, 'date': s.date, 'period': s.period, 'content': content,
            'notes': None, 'model': None, 'edited': False, 'class_name': ctx.cls.name, 'subject': ctx.ta.subject,
            'weekday': DAYS_FULL[s.date.weekday()], 'warnings': [], 'meta': content['_meta']}


class HeaderIn(BaseModel):
    school: str = Field(min_length=5, max_length=300)


@router.put('/daily-plans-header')
def put_header(body: HeaderIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    """Bütün gündəlik planların başlığında məktəbin adı (müəllimin öz seçimi)."""
    user.ai_settings = {**(user.ai_settings or {}), 'header_school': body.school.strip()}
    audit(db, user, 'update', 'daily_plan_header', user.id)
    db.commit()
    return {'school': header_school(user)}


class SlotIn(BaseModel):
    date: dt.date
    period: int


@router.post('/daily-plans/{ta_id}/manual')
def manual(ta_id: int, body: SlotIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    """Süni intellektsiz: şablondan (perspektiv plandan doldurulmuş) plan yaradılır, müəllim hər sətri özü yazır."""
    ctx = plan_ctx(db, own_assignment(db, user, ta_id))
    s = _slot(ctx, body.date, body.period)
    if db.scalar(select(DailyPlan.id).where(DailyPlan.assignment_id == ctx.ta.id, DailyPlan.date == s.date,
                                            DailyPlan.period == s.period)):
        raise HTTPException(409, 'Bu dərs üçün gündəlik plan artıq var')
    b = _blank(db, user, ctx, s)
    entry = db.scalar(select(JournalEntry).where(JournalEntry.assignment_id == ctx.ta.id, JournalEntry.date == s.date,
                                                 JournalEntry.period == s.period))
    pl = taught_lesson(ctx, s, entry)
    p = DailyPlan(assignment_id=ctx.ta.id, date=s.date, period=s.period, created_by=user.id,
                  plan_lesson_id=pl.id if pl else None, topic=b['topic'], content=b['content'], provider='manual',
                  model=None, edited=True)
    db.add(p)
    db.flush()
    audit(db, user, 'create', 'daily_plan', p.id, date=str(s.date), period=s.period, model='manual')
    db.commit()
    return {**_full(p, ctx), 'blank': False}


@router.get('/daily-plans/{ta_id}/preview')
def preview(ta_id: int, date: dt.date, period: int, user: User = Depends(staff), db: Session = Depends(get_db)):
    """Sistemdə baxış: hazır plan varsa – o, yoxdursa perspektiv plandan doldurulmuş şablon."""
    ctx = plan_ctx(db, own_assignment(db, user, ta_id))
    s = _slot(ctx, date, period)
    p = db.scalar(select(DailyPlan).where(DailyPlan.assignment_id == ctx.ta.id, DailyPlan.date == date,
                                          DailyPlan.period == period))
    return {**_full(p, ctx), 'blank': False} if p else _blank(db, user, ctx, s)


class GenIn(BaseModel):
    date: dt.date
    period: int
    notes: str | None = Field(None, max_length=1500)


@router.post('/daily-plans/{ta_id}/prompt')
def prompt_preview(ta_id: int, body: GenIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    """Promtun özünü göstərir (şəffaflıq: müəllim hansı məlumatın göndərildiyini görür). Açar tələb olunmur."""
    ctx = plan_ctx(db, own_assignment(db, user, ta_id))
    c, _ = _context(db, user, ctx, _slot(ctx, body.date, body.period), body.notes)
    return {'system': dp.system_prompt(c['minutes']), 'user': dp.user_prompt(c)}


@router.post('/daily-plans/{ta_id}/generate')
def generate(ta_id: int, body: GenIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    ctx = plan_ctx(db, own_assignment(db, user, ta_id))
    s = _slot(ctx, body.date, body.period)
    if db.scalar(select(DailyPlan.id).where(DailyPlan.assignment_id == ctx.ta.id, DailyPlan.date == s.date,
                                            DailyPlan.period == s.period)):
        raise HTTPException(409, 'Bu dərs üçün gündəlik plan artıq hazırdır – yenisini hazırlamaq üçün əvvəlcə onu silin')
    cfg = ai_config(user)
    c, pl = _context(db, user, ctx, s, body.notes)
    raw = ai.complete_json(cfg, dp.system_prompt(c['minutes']), dp.user_prompt(c))
    content, warn = dp.normalize(raw, c)
    content['_warnings'] = warn
    content['_meta'] = {k: c[k] for k in META_KEYS}
    p = DailyPlan(assignment_id=ctx.ta.id, date=s.date, period=s.period, created_by=user.id)   # hər dərs saatı – ayrıca plan
    db.add(p)
    p.plan_lesson_id, p.topic, p.content, p.notes = pl.id, pl.topic, content, c['notes']
    p.provider, p.model, p.edited = cfg['provider'], cfg['model'], False
    db.flush()
    audit(db, user, 'create', 'daily_plan', p.id, date=str(s.date), period=s.period, model=cfg['model'])
    db.commit()
    return _full(p, ctx)


def _own_plan(db: Session, user: User, ta_id: int, pid: int):
    ctx = plan_ctx(db, own_assignment(db, user, ta_id))
    p = db.get(DailyPlan, pid)
    if not p or p.assignment_id != ctx.ta.id:
        raise HTTPException(404, 'Gündəlik plan tapılmadı')
    return p, ctx


@router.get('/daily-plans/{ta_id}/item/{pid}')
def get_plan(ta_id: int, pid: int, user: User = Depends(staff), db: Session = Depends(get_db)):
    p, ctx = _own_plan(db, user, ta_id, pid)
    return _full(p, ctx)


class EditIn(BaseModel):
    content: dict


@router.put('/daily-plans/{ta_id}/item/{pid}')
def edit_plan(ta_id: int, pid: int, body: EditIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    p, ctx = _own_plan(db, user, ta_id, pid)
    keep = {k: v for k, v in (p.content or {}).items() if k.startswith('_')}
    new = {k: v for k, v in body.content.items() if not k.startswith('_')}
    total = sum(int(m.get('vaxt') or 0) for m in new.get('merheleler') or [] if isinstance(m, dict))
    minutes = (keep.get('_meta') or {}).get('minutes', dp.LESSON_MIN)
    keep['_warnings'] = [] if total == minutes else [f'Mərhələlərin vaxtı cəmi {total} dəq (olmalı: {minutes}) – yoxlayın']
    p.content = {**new, **keep}
    p.edited = True
    if isinstance(new.get('basliq'), dict) and str(new['basliq'].get('topic') or '').strip():
        p.topic = str(new['basliq']['topic']).strip()[:1000]
    audit(db, user, 'update', 'daily_plan', p.id)
    db.commit()
    return _full(p, ctx)


@router.delete('/daily-plans/{ta_id}/item/{pid}')
def delete_plan(ta_id: int, pid: int, user: User = Depends(staff), db: Session = Depends(get_db)):
    p, _ = _own_plan(db, user, ta_id, pid)
    db.delete(p)
    audit(db, user, 'delete', 'daily_plan', pid)
    db.commit()
    return {'ok': True}


# ---------------------------------------------------------------- Word (.docx)
@router.get('/daily-plans-docx')
def docx_slots(slots: str, user: User = Depends(staff), db: Session = Depends(get_db)):
    """slots = «ta:YYYY-MM-DD:saat,…» – hazır plan, yoxdursa şablon; bir Word faylında, verilən sıra ilə."""
    docs, ctxs = [], {}
    for part in slots.split(',')[:60]:
        try:
            t, d, per = part.split(':')
            t, d, per = int(t), dt.date.fromisoformat(d), int(per)
        except ValueError:
            continue
        if t not in ctxs:
            ctxs[t] = plan_ctx(db, own_assignment(db, user, t))
        ctx = ctxs[t]
        s = slot_at(ctx.slots, d, per)
        if not s:
            continue
        p = db.scalar(select(DailyPlan).where(DailyPlan.assignment_id == t, DailyPlan.date == d, DailyPlan.period == per))
        if p:
            docs.append((p.topic, p.content or {}))
            continue
        try:
            b = _blank(db, user, ctx, s)
        except HTTPException:
            continue
        docs.append((b['topic'], b['content']))
    if not docs:
        raise HTTPException(404, 'Perspektiv planda mövzusu olan dərs tapılmadı')
    buf = io.BytesIO()
    build_docx(docs, header_school(user)).save(buf)
    buf.seek(0)
    from urllib.parse import quote
    name = f'Gundelik planlasdirma {dt.date.today():%d.%m.%Y}.docx'
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                             headers={'Content-Disposition': f"attachment; filename*=UTF-8''{quote(name)}"})


def build_docx(docs: list[tuple[str, dict]], school: str = SCHOOL_DEFAULT):
    """ARTİ «Gündəlik planlaşdırma» forması: Məktəb/Fənn/Müəllim/Sinif/Tarix, Altstandart(lar), Təlim nəticəsi(ləri),
    Qiymətləndirmə meyar(lar)ı, Mövzu, Dərsin təşkili, İş üsulu / İş forması, Refleksiya. Boş bölmə – yazmaq üçün xətlər."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Mm(18))
    st = doc.styles['Normal']
    st.font.name, st.font.size = 'Times New Roman', Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(2)

    def para(text='', bold=False, size=None, center=False, before=0):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold = bold
            if size:
                r.font.size = Pt(size)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        return p

    def labeled(label, value, before=0):
        p = para(before=before)
        p.add_run(label).bold = True
        p.add_run(value)
        return p

    def field(p, label, value):
        p.add_run(label + ': ').bold = True
        p.add_run(value or '____________________')

    def lines(n=3):
        for _ in range(n):
            para('_' * 78)

    def heading(t):
        para(t, bold=True, before=10)

    def no_borders(t):
        b = OxmlElement('w:tblBorders')
        for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement(f'w:{e}')
            el.set(qn('w:val'), 'nil')
            b.append(el)
        t._tbl.tblPr.append(b)

    for n, (topic, c) in enumerate(docs):
        h = {k: v for k, v in (c.get('basliq') or {}).items() if isinstance(v, str) and v.strip()}
        m = {**c.get('_meta', {}), 'school': school, **h}             # redaktə olunmuş başlıq sətirləri üstündür
        topic = h.get('topic', topic)
        if n:
            doc.add_page_break()
        para('Gündəlik planlaşdırma', bold=True, size=15, center=True)
        t = doc.add_table(rows=3, cols=2)
        no_borders(t)
        field(t.cell(0, 0).paragraphs[0], 'Məktəb', m.get('school', ''))
        field(t.cell(0, 1).paragraphs[0], 'Fənn', m.get('subject', ''))
        field(t.cell(1, 0).paragraphs[0], 'Müəllim', m.get('teacher', ''))
        field(t.cell(1, 1).paragraphs[0], 'Sinif', m.get('class_name', ''))
        field(t.cell(2, 1).paragraphs[0], 'Tarix', m.get('tarix') or f"{m.get('date_text', '')} ({m.get('period', '')}-ci dərs)")
        for row in t.rows:
            row.cells[0].width, row.cells[1].width = Mm(108), Mm(66)

        heading('Altstandart(lar):')
        std = c.get('standartlar') or []
        for s in std:
            para(f"{s.get('kod', '')}{(' – ' + s['metn']) if s.get('metn') else ''}")
        if not any(s.get('metn') for s in std):
            lines(2)
        heading('Təlim nəticəsi(ləri):')
        if c.get('telim_neticeleri'):
            for i, x in enumerate(c['telim_neticeleri'], 1):
                para(f'{i}. {x}')
        else:
            lines(3)
        heading('Qiymətləndirmə meyar(lar)ı:')
        q = c.get('qiymetlendirme') or {}
        if q.get('meyarlar'):
            for x in q['meyarlar']:
                para(f'• {x}')
        else:
            lines(3)
        if q.get('spesifikasiya'):
            tt = doc.add_table(rows=1, cols=4)
            tt.style = 'Table Grid'
            for cell, h in zip(tt.rows[0].cells, ['Altstandart', 'Tapşırıq sayı', 'Çətinlik', 'Bal']):
                cell.text = h
            for sp in q['spesifikasiya']:
                r = tt.add_row().cells
                for cell, k in zip(r, ('standart', 'tapsiriq_sayi', 'seviyye', 'bal')):
                    cell.text = str(sp.get(k, ''))
        labeled('Mövzu: ', topic, before=10)
        heading('Dərsin təşkili (Şagirdlərin dərsə cəlbolunması, sual və tapşırıqlar):')
        stages = c.get('merheleler') or []
        if c.get('tedqiqat_suali'):
            labeled('Tədqiqat sualı: ', c['tedqiqat_suali'])
        for i, s in enumerate(stages, 1):
            para(f"{i}. {s.get('ad', '')}" + (f" ({s['vaxt']} dəq)" if s.get('vaxt') else ''), bold=True, before=4)
            if s.get('muellim'):
                labeled('Müəllim: ', s['muellim'])
            if s.get('sagird'):
                labeled('Şagirdlər: ', s['sagird'])
            for j, x in enumerate(s.get('tapsiriqlar') or [], 1):
                para(f"    {j}) {x.get('metn', '')}" + (f"  [Cavab: {x['cavab']}]" if x.get('cavab') else ''))
        d = c.get('diferensial') or {}
        if d.get('destek'):
            labeled('Dəstək (zəif şagirdlər): ', d['destek'], before=4)
        if d.get('inkisaf'):
            labeled('İnkişaf (güclü şagirdlər): ', d['inkisaf'])
        if not stages:
            lines(12)
        if c.get('ev_tapsirigi'):
            labeled('Ev tapşırığı: ', c['ev_tapsirigi'], before=4)
        t = doc.add_table(rows=1, cols=2)
        no_borders(t)
        for cell, label, vals in ((t.cell(0, 0), 'İş üsulu:', c.get('is_usullari')),
                                  (t.cell(0, 1), 'İş forması:', c.get('is_formalari'))):
            cell.paragraphs[0].add_run(label).bold = True
            cell.paragraphs[0].paragraph_format.space_before = Pt(10)
            for x in vals or ['_' * 32] * 3:
                cell.add_paragraph(x)
        heading('Refleksiya:')
        if c.get('refleksiya'):
            for x in c['refleksiya']:
                para(f'• {x}')
        else:
            lines(3)
    return doc
