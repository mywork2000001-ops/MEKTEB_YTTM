"""Analitika və hesabat: icmal, reytinq (ümumi və irəliləyiş), güclü/orta/zəif (avtomatik), risk,
davamiyyət xəritəsi (25%+), şagird kartı (valideynlə əlaqə, fərdi iş planı), Excel ixracı, audit jurnalı."""
from __future__ import annotations

import datetime as dt
import io
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..analytics import Period, analyze, attendance_map
from ..db import get_db
from ..deps import admin_only, staff
from ..models import AuditLog, IndividualPlan, ParentContact, Student, TeachingAssignment, User
from ..services import own_assignment, plan_ctx
from .common import audit, can_see_student, get_or_404

router = APIRouter(prefix='/api', tags=['analytics'])


def _period(ctx, date_from, date_to, semester) -> Period:
    y = ctx.year
    if semester == 1:
        return Period(y.start, y.sem1_end)
    if semester == 2:
        return Period(y.sem2_start, y.end)
    a, b = date_from or y.start, date_to or y.end
    if b < a:
        raise HTTPException(400, 'Dövrün sonu əvvəlindən qabaq ola bilməz')
    return Period(a, b)


@router.get('/analytics/{ta_id}')
def class_analytics(ta_id: int, date_from: dt.date | None = None, date_to: dt.date | None = None,
                    semester: int | None = None, user: User = Depends(staff), db: Session = Depends(get_db)):
    ctx = plan_ctx(db, own_assignment(db, user, ta_id))
    return analyze(db, ctx, _period(ctx, date_from, date_to, semester))


@router.get('/analytics/{ta_id}/levels')
def levels(ta_id: int, semester: int | None = None, user: User = Depends(staff), db: Session = Depends(get_db)):
    """Güclü / orta / zəif – nəticələrə görə avtomatik (nəticə yoxdursa IX sinif balı)."""
    ctx = plan_ctx(db, own_assignment(db, user, ta_id))
    rows = analyze(db, ctx, _period(ctx, None, None, semester))['students']
    out = {k: [{'student_id': r['student_id'], 'full_name': r['full_name'], 'rating': r['rating'],
                'source': r['level_source']} for r in rows if r['level'] == k] for k in ('Güclü', 'Orta', 'Zəif')}
    out['Məlum deyil'] = [{'student_id': r['student_id'], 'full_name': r['full_name'], 'rating': None, 'source': None}
                          for r in rows if r['level'] is None]          # nə nəticə, nə IX balı var
    return out


@router.get('/analytics/{ta_id}/attendance')
def attendance(ta_id: int, date_from: dt.date | None = None, date_to: dt.date | None = None,
               semester: int | None = None, user: User = Depends(staff), db: Session = Depends(get_db)):
    ctx = plan_ctx(db, own_assignment(db, user, ta_id))
    return attendance_map(db, ctx, _period(ctx, date_from, date_to, semester))


@router.get('/overview')
def overview(user: User = Depends(staff), db: Session = Depends(get_db)):
    """Bütün siniflərim üzrə qısa icmal (Əsas səhifə / Analitika)."""
    out = []
    for ta in db.scalars(select(TeachingAssignment).where(TeachingAssignment.teacher_id == user.id,
                                                          TeachingAssignment.archived_at.is_(None))):
        ctx = plan_ctx(db, ta)
        if ctx.cls.archived_at:
            continue
        a = analyze(db, ctx, _period(ctx, None, None, None), links=False)
        out.append({'ta_id': ta.id, 'class_name': ctx.cls.name, 'subject': ta.subject, **a['overview'],
                    'lessons_written': a['lessons_written']})
    return sorted(out, key=lambda x: x['class_name'])


@router.get('/my/weekly')
def weekly(user: User = Depends(staff), db: Session = Depends(get_db)):
    """Həftəlik xülasə (Əsas səhifə; serverdə saxlanmır – hər dəfə hesablanır): siniflərim üzrə qırmızı risk,
    25%-dən çox buraxanlar, son 7 gündə yazılmamış dərslər. week – ISO həftə (cihazda «oxundu» açarı)."""
    from ..performance import lesson_counts
    from ..services import today
    t = today()
    out = []
    for ta in db.scalars(select(TeachingAssignment).where(TeachingAssignment.teacher_id == user.id,
                                                          TeachingAssignment.archived_at.is_(None))):
        ctx = plan_ctx(db, ta)
        if ctx.cls.archived_at:
            continue
        sem = 2 if t >= ctx.year.sem2_start else 1
        a = analyze(db, ctx, _period(ctx, None, None, sem), links=False)
        lc = lesson_counts(db, ctx, t, sem)
        recent = [m for m in lc['missing_list'] if (t - m['date']).days <= 7]
        red = [r['full_name'] for r in a['students'] if r['risk']['status'] == 'Qırmızı']
        absent = [r['full_name'] for r in a['students'] if r['absence_warning']]
        if red or absent or recent:
            out.append({'ta_id': ta.id, 'class_name': ctx.cls.name, 'subject': ta.subject, 'red': red,
                        'absence': absent, 'missing_week': recent, 'missing_total': lc['missing']})
    y, w, _ = t.isocalendar()
    return {'week': f'{y}-W{w:02d}', 'today': t, 'items': sorted(out, key=lambda x: x['class_name'])}


class LevelIn(BaseModel):
    level: Literal['Güclü', 'Orta', 'Zəif'] | None = None     # None – avtomatik
    note: str | None = Field(None, max_length=300)


@router.put('/analytics/{ta_id}/levels/{sid}')
def set_level(ta_id: int, sid: int, body: LevelIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    """Müəllim şagirdin səviyyəsini əl ilə təyin edir (öz fənni üzrə); None – yenidən avtomatik."""
    from ..services import roster
    ta = own_assignment(db, user, ta_id)
    if sid not in {s.id for s in roster(db, ta)}:
        raise HTTPException(404, 'Şagird bu sinifdə/qrupda deyil')
    from .levels import _set
    _set(db, user, ta.id, sid, body.level, 'manual', True, note=body.note)   # müəllimin qərarı – kilidli, tarixçədə
    audit(db, user, 'update', 'level', sid, level=body.level)
    db.commit()
    return {'ok': True, 'level': body.level}


# ---------------------------------------------------------------- şagird kartı (müəllimin öz qeydləri)
def _student(db: Session, user: User, sid: int) -> Student:
    s = get_or_404(db, Student, sid, 'Şagird')
    if not can_see_student(db, user, s):
        raise HTTPException(403, 'Bu şagird sizin siniflərinizdən deyil')
    return s


class ContactIn(BaseModel):
    date: dt.date
    method: Literal['zəng', 'görüş', 'mesaj', 'iclas']
    topic: str = Field(min_length=2, max_length=300)
    outcome: str | None = Field(None, max_length=3000)
    follow_up: dt.date | None = None


def contact_out(c: ParentContact):
    return {'id': c.id, 'date': c.date, 'method': c.method, 'topic': c.topic, 'outcome': c.outcome,
            'follow_up': c.follow_up}


@router.get('/students/{sid}/contacts')
def contacts(sid: int, user: User = Depends(staff), db: Session = Depends(get_db)):
    _student(db, user, sid)
    return [contact_out(c) for c in db.scalars(select(ParentContact).where(
        ParentContact.student_id == sid, ParentContact.teacher_id == user.id).order_by(ParentContact.date.desc()))]


@router.post('/students/{sid}/contacts')
def add_contact(sid: int, body: ContactIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    _student(db, user, sid)
    c = ParentContact(student_id=sid, teacher_id=user.id, **body.model_dump())
    db.add(c)
    db.flush()
    audit(db, user, 'create', 'parent_contact', c.id, student_id=sid)
    db.commit()
    return contact_out(c)


@router.delete('/students/{sid}/contacts/{cid}')
def del_contact(sid: int, cid: int, user: User = Depends(staff), db: Session = Depends(get_db)):
    c = db.get(ParentContact, cid)
    if not c or c.student_id != sid or c.teacher_id != user.id:
        raise HTTPException(404, 'Qeyd tapılmadı')
    db.delete(c)
    audit(db, user, 'delete', 'parent_contact', cid)
    db.commit()
    return {'ok': True}


class StepIn(BaseModel):
    text: str = Field(min_length=1, max_length=500)
    done: bool = False


class PlanIn(BaseModel):
    goal: str = Field(min_length=3, max_length=2000)
    steps: list[StepIn] = Field(default_factory=list)
    start: dt.date
    review_date: dt.date | None = None
    status: Literal['aktiv', 'tamamlandı', 'dayandırıldı'] = 'aktiv'
    note: str | None = Field(None, max_length=3000)


def iplan_out(p: IndividualPlan):
    return {'id': p.id, 'goal': p.goal, 'steps': p.steps, 'start': p.start, 'review_date': p.review_date,
            'status': p.status, 'note': p.note,
            'progress': round(sum(s['done'] for s in p.steps) * 100 / len(p.steps)) if p.steps else 0}


@router.get('/students/{sid}/plans')
def iplans(sid: int, user: User = Depends(staff), db: Session = Depends(get_db)):
    _student(db, user, sid)
    return [iplan_out(p) for p in db.scalars(select(IndividualPlan).where(
        IndividualPlan.student_id == sid, IndividualPlan.teacher_id == user.id).order_by(IndividualPlan.start.desc()))]


@router.post('/students/{sid}/plans')
def add_iplan(sid: int, body: PlanIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    _student(db, user, sid)
    p = IndividualPlan(student_id=sid, teacher_id=user.id, **{**body.model_dump(), 'steps': [s.model_dump() for s in body.steps]})
    db.add(p)
    db.flush()
    audit(db, user, 'create', 'individual_plan', p.id, student_id=sid)
    db.commit()
    return iplan_out(p)


@router.put('/students/{sid}/plans/{pid}')
def upd_iplan(sid: int, pid: int, body: PlanIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    p = db.get(IndividualPlan, pid)
    if not p or p.student_id != sid or p.teacher_id != user.id:
        raise HTTPException(404, 'Plan tapılmadı')
    for k, v in body.model_dump().items():
        setattr(p, k, v)
    p.steps = [s.model_dump() for s in body.steps]
    audit(db, user, 'update', 'individual_plan', p.id)
    db.commit()
    return iplan_out(p)


# ---------------------------------------------------------------- Excel ixracı
@router.get('/reports/{ta_id}/xlsx')
def export_xlsx(ta_id: int, semester: int | None = None, user: User = Depends(staff), db: Session = Depends(get_db)):
    """Sinif hesabatı (ağ-qara çap üçün, A4): vərəqlər – Reytinq, Müvəffəqiyyət, Davamiyyət, Dərs sayı.
    Hər vərəqdə məktəb adı, sənədin adı, dövr və imza sətirləri; başlıq sətri hər səhifədə təkrarlanır."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from ..analytics import attendance_map
    from ..models import School
    from ..performance import MIN_MARKS, lesson_counts, metrics, subject_grades
    from ..services import roster, today
    from .lessonplans import SCHOOL_DEFAULT
    ctx = plan_ctx(db, own_assignment(db, user, ta_id))
    p = _period(ctx, None, None, semester)
    a = analyze(db, ctx, p)
    sc = db.get(School, user.school_id) if user.school_id else None
    school = sc.name if sc and sc.name else SCHOOL_DEFAULT
    deputy = ((sc.doc_settings or {}) if sc else {}).get('deputy') or ''
    sem_label = f'{semester}-ci yarımil' if semester in (1, 2) else 'bütün il'
    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    grey = PatternFill('solid', fgColor='D9D9D9')                 # Canon ağ-qara: rəng yox, yalnız boz fon
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title: str, doc: str, head: list[str], rows: list[list], widths: list[int], notes: tuple[str, ...] = ()):
        ws = wb.create_sheet(title)
        n = len(head)
        lines = [(school, True, 11), (doc, True, 12),
                 (f'{a["class_name"]} · {a["subject"]} · {sem_label}: {p.a:%d.%m.%Y} – {p.b:%d.%m.%Y} · '
                  f'müəllim: {user.full_name}', False, 10)]
        for i, (txt, bold, size) in enumerate(lines, 1):
            ws.cell(i, 1, txt).font = Font(bold=bold, size=size)
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=n)
            ws.cell(i, 1).alignment = Alignment(horizontal='center')
        hr = 5
        for j, h in enumerate(head, 1):
            c = ws.cell(hr, j, h)
            c.font, c.fill, c.border = Font(bold=True), grey, border
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for r in rows:
            ws.append(r)
            for c in ws[ws.max_row]:
                c.border = border
        last = ws.max_row + 2
        for t in notes:
            ws.cell(last, 1, t).font = Font(italic=True, size=9)
            last += 1
        last += 1
        for role, name in (('Fənn müəllimi', user.full_name), ('Direktor müavini (tədris işləri üzrə)', deputy)):
            ws.cell(last, 1, f'{role}:')
            ws.cell(last, 3, '______________ (imza)')
            ws.cell(last, 5, name or '______________')
            last += 2
        ws.cell(last, 1, f'Tərtib edildi: {today():%d.%m.%Y}').font = Font(size=9)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f'{hr}:{hr}'
        ws.oddFooter.right.text = 'Səhifə &P / &N'
        return ws

    # 1. Reytinq
    ws = sheet('Reytinq', 'Şagirdlərin reytinqi',
               ['Yer', 'Şagird', 'IX bal', 'Formativ orta', 'KSQ orta %', 'Onlayn %', 'Ev tapşırığı %', 'Davamiyyət %',
                'Sınaq %', 'Reytinq', 'İrəliləyiş', 'Səviyyə'],
               [[r['place'], r['full_name'], r['ix_score'], r['avg_grade'], r['ksq_avg_pct'], r['online_pct'], r['homework_pct'],
                 r['attendance_pct'], r['sinaq_pct'], r['rating'], r['progress'], r['level']] for r in a['students']],
               [5, 32, 7, 9, 9, 9, 10, 11, 8, 8, 10, 9],
               (f'IX bal – {a["ix_label"]}. Reytinq: qiymət 40% · KSQ 30% · ev tapşırığı 15% · davamiyyət 15%; sınaq daxil deyil.',
                'Qalın davamiyyət – dərslərin 25%-dən çoxunu buraxıb.'))
    for i, r in enumerate(a['students'], 6):
        if r['absence_warning']:
            ws.cell(i, 8).font = Font(bold=True)
    # 2. Müvəffəqiyyət
    studs = roster(db, ctx.ta)
    g = subject_grades(db, ctx, semester if semester in (1, 2) else None, [s.id for s in studs])
    m = metrics([g[s.id]['grade'] for s in studs])
    sheet('Müvəffəqiyyət', 'Müvəffəqiyyət hesabatı',
          ['№', 'Şagird', 'Qiymət', 'Mənbə', 'Formativ orta', 'Qiymət sayı', 'Yarımil'],
          [[i, s.full_name, g[s.id]['grade'], g[s.id]['source'] or 'qiymət yoxdur', g[s.id]['formative_avg'], g[s.id]['marks'],
            g[s.id]['semester_grade']] for i, s in enumerate(studs, 1)],
          [5, 32, 8, 14, 11, 10, 9],
          (f'Şagird: {m["students"]} · qiymətləndirilib: {m["graded"]} · «5» {m["distribution"][5]} · «4» {m["distribution"][4]} · '
           f'«3» {m["distribution"][3]} · «2» {m["distribution"][2]}',
           f'Müvəffəqiyyət: {_pct(m["success_pct"])} · keyfiyyət: {_pct(m["quality_pct"])} · orta qiymət: {m["avg"] or "—"} · '
           f'SOU: {_pct(m["sou"])}',
           f'Qiymət: yarımil qiyməti (KSQ×0,4 + BSQ×0,6), yoxdursa formativ orta (ən azı {MIN_MARKS} qiymət).'))
    # 3. Davamiyyət
    am = attendance_map(db, ctx, p)
    sheet('Davamiyyət', 'Davamiyyət',
          ['Şagird', 'Qeyd olunan dərs', 'Buraxıb', 'O cümlədən qayıb', 'Gecikmə', 'Buraxma %'],
          [[r['full_name'], sum(1 for c in r['cells'] if c), r['missed'], r['unexcused'], r['late'], r['missed_pct']]
           for r in am['rows']],
          [32, 12, 10, 12, 10, 11],
          (f'Hədd: {am["limit_pct"]:g}% · yalnız bu fənnin jurnalda yazılmış dərsləri.',))
    # 4. Dərs sayı
    t = today()
    rows = [('I yarımil', lesson_counts(db, ctx, t, 1)), ('II yarımil', lesson_counts(db, ctx, t, 2)),
            ('Bütün il', lesson_counts(db, ctx, t))]
    sheet('Dərs sayı', f'Dərs sayı ({t:%d.%m.%Y} vəziyyəti)',
          ['Dövr', 'Həftədə', 'Planda', 'Cədvəldə', 'Keçilməli idi', 'Yazılıb', 'Yazılmayıb', 'Keçilən mövzu', 'Qalan', 'Geriləmə'],
          [[n, x['weekly_hours'], x['plan_total'], x['timetable_total'], x['due'], x['written'], x['missing'], x['covered'],
            x['remaining'], x['lag']] for n, x in rows],
          [14, 9, 9, 10, 12, 9, 11, 13, 8, 10])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f'hesabat_{ctx.cls.code}_{p.a:%Y%m%d}-{p.b:%Y%m%d}.xlsx'
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             headers={'Content-Disposition': f'attachment; filename="{name}"'})


def _pct(v) -> str:
    return '—' if v is None else f'{v:.1f}%'.replace('.', ',')


# ---------------------------------------------------------------- audit jurnalı (admin)
@router.get('/audit')
def audit_log(user_id: int | None = None, entity: str | None = None, limit: int = 200,
              user: User = Depends(admin_only), db: Session = Depends(get_db)):
    """Kim, nə vaxt, nəyi dəyişib. Mesajların məzmunu burada YOXDUR (məxfilik)."""
    st = select(AuditLog, User.full_name).join(User, User.id == AuditLog.user_id, isouter=True) \
        .where((User.school_id == user.school_id) | (AuditLog.user_id.is_(None)))
    if user_id:
        st = st.where(AuditLog.user_id == user_id)
    if entity:
        st = st.where(AuditLog.entity == entity)
    rows = db.execute(st.order_by(AuditLog.id.desc()).limit(min(limit, 1000)))
    return [{'id': a.id, 'at': a.at, 'user': n, 'action': a.action, 'entity': a.entity, 'entity_id': a.entity_id,
             'details': a.details} for a, n in rows]
