"""İşlənməyənlər siyahısının çap formaları (docs/tasks/islenmeyenler-siyahisi.md): sinif üzrə (A4, PDF/Excel, yalnız AZ) və
hər şagirdə ayrıca vərəqə «İşləməli olduğun testlər» (PDF; şagirdin dili ru isə AZ|RU). Üslub – class_matrix_print.py."""
from __future__ import annotations

import io
from html import escape as _e

from .class_matrix_print import KIND_LABEL, _d, period_text

DOC = 'İşlənməyən mövzu testləri və sınaqlar – sinif üzrə siyahı'
STUDENT_DOC = ('İşləməli olduğun testlər', 'Тесты, которые нужно выполнить')
CSS = ('@page{size:A4;margin:12mm}body{font:12px/1.4 "Times New Roman","Liberation Serif",serif;color:#000}'
       'h1{font-size:15px;text-align:center;margin:0 0 4px}p{margin:3px 0}.c{text-align:center}'
       'table{border-collapse:collapse;width:100%;margin-top:8px}th,td{border:1px solid #000;padding:3px 5px;vertical-align:top}'
       'th{background:#D9D9D9}thead{display:table-header-group}tr{page-break-inside:avoid}.sign{margin-top:24px}')


def item_text(it: dict) -> str:
    return ((it['no'] + ' ') if it.get('no') else '') + (it['title'] or '')


def _split(row: dict) -> tuple[list[str], list[str]]:
    mv = [item_text(i) for i in row['items'] if i['kind'] == 'movzu']
    other = [item_text(i) for i in row['items'] if i['kind'] != 'movzu']
    return mv, other


def class_html(m: dict, school: str, teacher: str) -> str:
    body = []
    for n, r in enumerate(m['rows'], 1):
        mv, other = _split(r)
        body.append(f"<tr><td class=\"c\">{n}</td><td>{_e(r['full_name'])}</td><td>{_e('; '.join(mv)) or '—'}</td>"
                    f"<td>{_e('; '.join(other)) or '—'}</td><td class=\"c\">{len(r['items'])}</td></tr>")
    if not body:
        body.append('<tr><td colspan="5" class="c">İşlənməyən test yoxdur.</td></tr>')
    return (
        '<!doctype html><html lang="az"><head><meta charset="utf-8"><title>' + _e(DOC) + f'</title><style>{CSS}</style></head><body>'
        f'<p class="c"><b>{_e(school)}</b></p><h1>{_e(DOC)}</h1>'
        f'<p class="c">{_e(m["class_name"])} · {_e(m["subject"])} · {_e(KIND_LABEL[m["kind"]])} · {_e(period_text(m))}'
        f' · müəllim: {_e(teacher)}</p>'
        '<table><thead><tr><th>№</th><th>Şagird</th><th>Mövzu testləri</th><th>Sınaqlar</th><th>Say</th></tr></thead><tbody>'
        + ''.join(body) +
        f'<tr><td></td><td><b>Cəmi: {m["students"]} şagird</b></td><td></td><td></td><td class="c"><b>{m["total"]}</b></td></tr>'
        '</tbody></table>'
        f'<p class="sign">Fənn müəllimi: {_e(teacher)} ______________ (imza)</p></body></html>')


def class_xlsx(m: dict, school: str, teacher: str, deputy: str = '') -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    grey = PatternFill('solid', fgColor='D9D9D9')
    wb = Workbook()
    ws = wb.active
    ws.title = 'İşlənməyənlər'
    for i, (txt, bold, size) in enumerate([(school, True, 11), (DOC, True, 12),
                                           (f'{m["class_name"]} · {m["subject"]} · {KIND_LABEL[m["kind"]]} · {period_text(m)} · müəllim: {teacher}', False, 10)], 1):
        ws.cell(i, 1, txt).font = Font(bold=bold, size=size)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
        ws.cell(i, 1).alignment = Alignment(horizontal='center')
    hr = 5
    for j, h in enumerate(['№', 'Şagird', 'Mövzu testləri', 'Sınaqlar', 'Say'], 1):
        c = ws.cell(hr, j, h)
        c.font, c.fill, c.border = Font(bold=True), grey, border
        c.alignment = Alignment(horizontal='center')

    def put(row: list):
        ws.append(row)
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center' if c.column in (1, 5) else 'left')

    for n, r in enumerate(m['rows'], 1):
        mv, other = _split(r)
        put([n, r['full_name'], '\n'.join(mv), '\n'.join(other), len(r['items'])])
    put(['', f'Cəmi: {m["students"]} şagird', '', '', m['total']])
    last = ws.max_row + 2
    for role, name in (('Fənn müəllimi', teacher), ('Direktor müavini (tədris işləri üzrə)', deputy)):
        ws.cell(last, 1, f'{role}:')
        ws.cell(last, 3, '______________ (imza)')
        ws.cell(last, 4, name or '______________')
        last += 2
    for col, w in zip('ABCDE', (5, 32, 48, 36, 7)):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{hr}:{hr}'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def students_html(m: dict, school: str, teacher: str) -> str:
    """Hər şagird ayrıca səhifə. lang=ru – başlıq və qeydlər AZ | RU."""
    pages = []
    for n, r in enumerate(m['rows']):
        ru = r.get('lang') == 'ru'

        def b(az: str, rus: str) -> str:
            return _e(az) + (f' / {_e(rus)}' if ru else '')
        trs = ''.join(f"<tr><td class=\"c\">{i}</td><td>{_e(item_text(it))}</td><td class=\"c\">{_e(_d(it['date']))}</td></tr>"
                      for i, it in enumerate(r['items'], 1))
        if not trs:
            trs = f'<tr><td colspan="3" class="c">{b("İşləməli test yoxdur.", "Нет тестов к выполнению.")}</td></tr>'
        brk = ' style="page-break-before:always"' if n else ''
        head = _e(STUDENT_DOC[0]) + (f' / {_e(STUDENT_DOC[1])}' if ru else '')
        pages.append(
            f'<section{brk}>'
            f'<p class="c"><b>{_e(school)}</b></p><h1>{head}</h1>'
            f'<p><b>{b("Şagird", "Ученик")}:</b> {_e(r["full_name"])}</p>'
            f'<p><b>{b("Sinif", "Класс")}:</b> {_e(r.get("class_name", m["class_name"]))} · <b>{b("Fənn", "Предмет")}:</b> {_e(r.get("subject", m["subject"]))}'
            f' · <b>{b("Dövr", "Период")}:</b> {_e(period_text(m))}</p>'
            f'<table><thead><tr><th>№</th><th>{b("Mövzu / sınaq", "Тема / пробный экзамен")}</th><th>{b("Tarix", "Дата")}</th></tr></thead>'
            f'<tbody>{trs}</tbody></table>'
            f'<p>{b("Portalda daxil ol → Testlər.", "Войдите в портал → Тесты.")}</p>'
            f'<p class="sign">{b("Fənn müəllimi", "Учитель")}: {_e(teacher)} ______________ ({b("imza", "подпись")})</p></section>')
    return '<!doctype html><html lang="az"><head><meta charset="utf-8"><title>' + _e(STUDENT_DOC[0]) + f'</title><style>{CSS}</style></head><body>' \
        + ''.join(pages) + '</body></html>'
